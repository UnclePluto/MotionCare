from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
import uuid

import openpyxl
import pytest
from django.db import connection, connections
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.patients.models import Patient
from apps.prescriptions.models import ActionLibraryItem
from apps.studies.models import ProjectPatient, StudyProject
from apps.training.export_rows import ExportRows
from apps.training.models import TrainingDetailExportLog, TrainingRecord
from apps.training.tests import test_export_rows as row_fixtures

from apps.training.tests.test_export_workbook import records
from apps.wearables.models import WearableMeasurement

export_sample = row_fixtures.export_sample
pytestmark = pytest.mark.django_db(transaction=True)


@pytest.fixture(autouse=True)
def action_library_for_transaction_tests(db):
    # transaction=True flushes migration seed data between cases.
    for key, name, kind in (
        ("game-executive-inhibition", "反应抑制", "game"),
        ("motion-resistance-shoulder-press", "肩推训练", "motion"),
    ):
        ActionLibraryItem.objects.get_or_create(
            source_key=key,
            defaults={
                "name": name,
                "internal_type": kind,
                "training_type": "游戏训练" if kind == "game" else "运动训练",
                "action_type": "训练",
                "has_ai_supervision": kind == "motion",
            },
        )


def login(user):
    client = APIClient(enforce_csrf_checks=True)
    client.get("/api/auth/csrf/")
    response = client.post(
        "/api/auth/login/",
        {"phone": user.phone, "password": "pass123456"},
        format="json",
        HTTP_X_CSRFTOKEN=client.cookies["csrftoken"].value,
    )
    assert response.status_code == 204
    client.credentials(HTTP_X_CSRFTOKEN=client.cookies["csrftoken"].value)
    return client


@pytest.fixture
def session_client(doctor):
    return login(doctor)


def download(client, pp, **payload):
    return client.post(
        f"/api/training/tracking/patients/{pp.patient_id}/export/",
        {"project_patient": pp.pk, "range": "all", **payload},
        format="json",
    )


def read_response(response):
    assert response.status_code == 200
    content = b"".join(response.streaming_content)
    response.close()
    return openpyxl.load_workbook(BytesIO(content))


def test_real_session_download_history_completed_project_and_audit(session_client, export_sample):
    pp = export_sample.project_patient
    StudyProject.objects.filter(pk=pp.project_id).update(status="completed")
    response = download(
        session_client, pp, range="custom", start_date="2026-09-09", end_date="2026-09-09"
    )
    assert response.status_code == 200
    assert (
        response["Content-Type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response["Cache-Control"] == "no-store"
    assert pp.patient.name not in unquote(response["Content-Disposition"])
    book = read_response(response)
    assert len(records(book["训练场次"])) == 4
    assert {r["处方版本"] for r in records(book["训练场次"])} == {1, 2}
    questions = records(book["游戏逐题"])
    assert len(questions) == 2
    assert {r["采集版本"] for r in questions} == {"active_response_v1"}
    assert "历史作答时间（毫秒）" not in questions[0]
    audit = TrainingDetailExportLog.objects.get()
    assert audit.status == "succeeded" and audit.error_code == ""
    assert audit.patient_id_snapshot == pp.patient_id
    assert audit.project_id_snapshot == pp.project_id
    assert audit.project_patient_id_snapshot == pp.pk
    assert audit.operator_id == pp.patient.primary_doctor_id
    assert audit.finished_at >= audit.started_at
    assert audit.format_version == "training_detail_v2"
    assert audit.filters == {
        "project_patient": pp.pk,
        "range": "custom",
        "start_date": "2026-09-09",
        "end_date": "2026-09-09",
    }
    assert audit.row_counts == {sheet.title: sheet.max_row - 1 for sheet in book}


def test_empty_export_and_invalid_filters(session_client, project_patient):
    book = read_response(download(session_client, project_patient))
    assert book.sheetnames == ["训练场次", "游戏逐题", "顺序选择明细", "运动明细", "训练期间生理数据", "字段说明"]
    assert all(s.max_row == 1 for s in list(book)[:4])
    assert (
        download(
            session_client,
            project_patient,
            range="custom",
            start_date="2026-09-10",
            end_date="2026-09-09",
        ).status_code
        == 400
    )
    assert TrainingDetailExportLog.objects.count() == 1


def test_no_session_missing_csrf_patient_role_no_audits(project_patient, doctor):
    assert download(APIClient(enforce_csrf_checks=True), project_patient).status_code == 403
    client = login(doctor)
    client.credentials()
    assert download(client, project_patient).status_code == 403
    user = User.objects.create_user(
        phone="13800002222", password="pass123456", name="患者账号", role="patient"
    )
    # Patient users can establish the same Django session; permissions still deny export.
    patient_client = APIClient(enforce_csrf_checks=True)
    assert patient_client.login(phone=user.phone, password="pass123456")
    patient_client.get("/api/auth/csrf/")
    patient_client.credentials(HTTP_X_CSRFTOKEN=patient_client.cookies["csrftoken"].value)
    assert download(patient_client, project_patient).status_code == 403
    assert TrainingDetailExportLog.objects.count() == 0


def test_row_scope_mismatch_and_unknown_are_same_404(session_client, project_patient, settings):
    settings.TRAINING_HEALTH_ENFORCE_ROW_SCOPE = True
    other = User.objects.create_user(phone="13800003333", password="pass123456", role="doctor")
    patient = Patient.objects.create(
        name="无权患者", gender="unknown", age=60, primary_doctor=other
    )
    project = StudyProject.objects.create(name="无权项目", created_by=other)
    pp = ProjectPatient.objects.create(patient=patient, project=project, created_by=other)
    responses = [
        download(session_client, pp),
        session_client.post(
            f"/api/training/tracking/patients/{project_patient.patient_id}/export/",
            {"project_patient": pp.pk},
            format="json",
        ),
        session_client.post(
            "/api/training/tracking/patients/999999/export/",
            {"project_patient": 999999},
            format="json",
        ),
    ]
    assert [r.status_code for r in responses] == [404, 404, 404]
    assert responses[0].data == responses[1].data == responses[2].data
    assert TrainingDetailExportLog.objects.count() == 0
    read_response(download(login(other), pp))


def test_audit_survives_real_unbind_and_operator_deletion(session_client, project_patient, doctor):
    read_response(download(session_client, project_patient))
    response = session_client.post(f"/api/studies/project-patients/{project_patient.pk}/unbind/")
    assert response.status_code == 200
    assert not ProjectPatient.objects.filter(pk=project_patient.pk).exists()
    audit = TrainingDetailExportLog.objects.get()
    assert audit.project_patient_id_snapshot == project_patient.pk
    doctor.delete()
    audit.refresh_from_db()
    assert audit.operator_id is None


def test_two_connections_keep_old_snapshot_until_next_export(
    session_client, export_sample, monkeypatch
):
    from apps.training import export_rows as source

    original = source._prepare_chunk
    changed = False

    def write_on_other_connection():
        try:
            TrainingRecord.objects.filter(pk=export_sample.revised.pk).update(
                motion_total_count=9, motion_standard_count=8
            )
            old = export_sample.measurements[0]
            WearableMeasurement.objects.create(
                provider=old.provider,
                patient_id=old.patient_id,
                device_id=old.device_id,
                metric_type="heart_rate",
                measured_at=export_sample.started + timedelta(seconds=1),
                attribution_status="attributed",
                source_fingerprint=str(uuid.uuid4()),
                heart_rate=90,
            )
        finally:
            connections.close_all()

    def modify_after_first_chunk(records, pp, rows):
        nonlocal changed
        original(records, pp, rows)
        if not changed:
            changed = True
            with ThreadPoolExecutor(max_workers=1) as pool:
                pool.submit(write_on_other_connection).result(timeout=10)

    monkeypatch.setattr(source, "RECORD_CHUNK_SIZE", 1)
    monkeypatch.setattr(source, "_prepare_chunk", modify_after_first_chunk)
    first = read_response(download(session_client, export_sample.project_patient))
    second = read_response(download(session_client, export_sample.project_patient))
    assert records(first["运动明细"])[0]["动作总次数"] == 5
    assert records(second["运动明细"])[0]["动作总次数"] == 9
    assert len(records(first["训练期间生理数据"])) == 8
    assert len(records(second["训练期间生理数据"])) == 10
    first_motion = next(
        r for r in records(first["训练场次"]) if r["训练记录编号"] == export_sample.revised.pk
    )
    assert first_motion["心率均值（次/分）"] == 70.5
    assert first_motion["心率读数数量"] == 2


@pytest.mark.parametrize(
    "failure,expected,code",
    [
        ("limit", 400, "limit_exceeded"),
        ("deadline", 503, "deadline_exceeded"),
        ("file", 500, "generation_failed"),
        ("audit", 500, "audit_failed"),
        ("response", 500, "generation_failed"),
    ],
)
def test_failures_audited_sanitized_and_resources_clean(
    session_client, export_sample, monkeypatch, failure, expected, code
):
    from apps.training import (
        export_views as views,
        export_workbook as engine,
        export_rows as source,
    )

    paths, handles = [], []
    old_rows_init = ExportRows.__init__

    def track_rows(self, *args, **kwargs):
        old_rows_init(self, *args, **kwargs)
        paths.append(self.directory)

    monkeypatch.setattr(ExportRows, "__init__", track_rows)
    original_temp = engine.TemporaryFile

    def track_temp(*args, **kwargs):
        handle = original_temp(*args, **kwargs)
        handles.append(handle)
        return handle

    monkeypatch.setattr(engine, "TemporaryFile", track_temp)
    if failure == "limit":
        monkeypatch.setattr(source, "MAX_RECORDS", 1)
    elif failure == "deadline":
        monkeypatch.setattr(views, "EXPORT_SECONDS", -1)
    elif failure == "file":
        from zipfile import ZipFile

        monkeypatch.setattr(
            ZipFile, "write", lambda *a, **kw: (_ for _ in ()).throw(OSError("SECRET_OBJECT/SQL"))
        )
    elif failure == "audit":
        original_save = TrainingDetailExportLog.save

        def fail_success(self, *args, **kwargs):
            if self.status == "succeeded":
                raise OSError("SECRET_AUDIT")
            return original_save(self, *args, **kwargs)

        monkeypatch.setattr(TrainingDetailExportLog, "save", fail_success)
    else:
        monkeypatch.setattr(
            views,
            "FileResponse",
            lambda *a, **kw: (_ for _ in ()).throw(OSError("SECRET_RESPONSE")),
        )
    response = download(session_client, export_sample.project_patient)
    assert response.status_code == expected
    assert "SECRET" not in str(response.data)
    if expected == 500:
        assert response.data == {"detail": "导出失败，请重试"}
    else:
        assert "缩小" in response.data["detail"]
    audit = TrainingDetailExportLog.objects.get()
    assert audit.status == "failed" and audit.error_code == code
    assert audit.finished_at is not None
    assert all(not Path(path).exists() for path in paths)
    assert all(handle.closed for handle in handles)


def test_success_response_close_releases_file_and_rows(session_client, export_sample, monkeypatch):
    from apps.training import export_workbook as engine

    handles = []
    original = engine.TemporaryFile

    def tracked(*args, **kwargs):
        handle = original(*args, **kwargs)
        handles.append(handle)
        return handle

    monkeypatch.setattr(engine, "TemporaryFile", tracked)
    response = download(session_client, export_sample.project_patient)
    assert response.status_code == 200 and not handles[0].closed
    response.close()
    assert handles[0].closed


def test_snapshot_transaction_read_only_and_first_sql(session_client, project_patient, monkeypatch):
    from apps.training import export_views as views

    original = views.prepare_export_rows
    captured = {}

    def inspect_snapshot(**kwargs):
        with connection.cursor() as cursor:
            cursor.execute("SHOW transaction_isolation")
            captured["isolation"] = cursor.fetchone()[0]
            cursor.execute("SHOW transaction_read_only")
            captured["read_only"] = cursor.fetchone()[0]
            cursor.execute("SELECT transaction_timestamp()")
            captured["at"] = cursor.fetchone()[0]
        return original(**kwargs)

    monkeypatch.setattr(views, "prepare_export_rows", inspect_snapshot)
    sql_in_transaction = []

    def capture_sql(execute, sql, params, many, context):
        if connection.in_atomic_block:
            sql_in_transaction.append(sql)
        return execute(sql, params, many, context)

    with connection.execute_wrapper(capture_sql):
        book = read_response(download(session_client, project_patient))
    assert sql_in_transaction[0] == "SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY"
    assert captured["isolation"] == "repeatable read" and captured["read_only"] == "on"
    metadata = {r["字段或项目"]: r["说明或值"] for r in records(book["字段说明"])}
    expected = captured["at"].astimezone(views.SHANGHAI).replace(tzinfo=None)
    assert abs((metadata["数据读取截至点"] - expected).total_seconds()) < 0.001


@pytest.mark.parametrize("server_cursor", [False, True])
def test_remaining_budget_cancels_real_sql_and_server_fetch(monkeypatch, server_cursor):
    from django.db import DatabaseError, transaction
    from apps.training import _export_deadline as budget

    clock = [100.0]
    monkeypatch.setattr(budget, "_monotonic", lambda: clock[0])
    with transaction.atomic():
        with connection.cursor() as cursor:
            cursor.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY")
        with budget.StatementBudget(deadline=160.0):
            cursor_factory = connection.chunked_cursor if server_cursor else connection.cursor
            with cursor_factory() as cursor:
                # DECLARE sees 60 seconds; FETCH must see only the remaining 40 ms.
                if server_cursor:
                    cursor.execute("SELECT pg_sleep(0.3)")
                clock[0] = 159.96
                with pytest.raises(DatabaseError) as error:
                    if server_cursor:
                        cursor.fetchmany(1)
                    else:
                        cursor.execute("SELECT pg_sleep(0.3)")
                assert getattr(error.value.__cause__, "sqlstate", None) == "57014"
        transaction.set_rollback(True)


def test_scope_revoked_between_authorizations_returns_404_and_failed_audit(
    session_client, project_patient, monkeypatch, settings
):
    settings.TRAINING_HEALTH_ENFORCE_ROW_SCOPE = True
    other = User.objects.create_user(phone="13800005555", password="pass123456", role="doctor")
    original_create = TrainingDetailExportLog.objects.create

    def revoke_after_audit(**kwargs):
        audit = original_create(**kwargs)
        Patient.objects.filter(pk=project_patient.patient_id).update(primary_doctor=other)
        StudyProject.objects.filter(pk=project_patient.project_id).update(created_by=other)
        return audit

    monkeypatch.setattr(TrainingDetailExportLog.objects, "create", revoke_after_audit)
    response = download(session_client, project_patient)
    assert response.status_code == 404
    audit = TrainingDetailExportLog.objects.get()
    assert audit.status == "failed" and audit.error_code == "scope_changed"


def test_audit_start_failure_is_fixed_json_no_workbook(
    session_client, project_patient, monkeypatch
):
    def fail(**kwargs):
        raise OSError("SECRET_AUDIT")

    monkeypatch.setattr(TrainingDetailExportLog.objects, "create", fail)
    response = download(session_client, project_patient)
    assert response.status_code == 500 and response.data == {"detail": "导出失败，请重试"}
    assert TrainingDetailExportLog.objects.count() == 0


@pytest.mark.parametrize("server_cursor", [False, True])
def test_api_real_database_cancel_is_503_audited_and_scope_is_restored(
    session_client, project_patient, monkeypatch, server_cursor
):
    from apps.training import export_views as views

    def slow_rows(**kwargs):
        factory = connection.chunked_cursor if server_cursor else connection.cursor
        with factory() as cursor:
            cursor.execute("SELECT pg_sleep(0.3)")
            if server_cursor:
                cursor.fetchmany(1)

    monkeypatch.setattr(views, "prepare_export_rows", slow_rows)
    monkeypatch.setattr(views, "EXPORT_SECONDS", 0.05)
    response = download(session_client, project_patient)
    assert response.status_code == 503 and "缩小" in response.data["detail"]
    audit = TrainingDetailExportLog.objects.get()
    assert audit.status == "failed" and audit.error_code == "deadline_exceeded"
    assert connection.execute_wrappers == []
    assert connection.get_autocommit()
    with connection.cursor() as cursor:
        cursor.execute("SHOW transaction_read_only")
        assert cursor.fetchone()[0] == "off"
        cursor.execute("SHOW statement_timeout")
        assert cursor.fetchone()[0] == "0"


def test_administrator_can_export_outside_doctor_scope(project_patient, settings):
    settings.TRAINING_HEALTH_ENFORCE_ROW_SCOPE = True
    admin = User.objects.create_user(phone="13800009999", password="pass123456", role="admin")
    book = read_response(download(login(admin), project_patient))
    assert book.sheetnames[0] == "训练场次"
    assert TrainingDetailExportLog.objects.get().operator_id == admin.pk


def test_valid_basic_credentials_without_session_or_csrf_cannot_export(
    doctor, project_patient, monkeypatch
):
    from base64 import b64encode
    from apps.training import export_views as views

    generated = []
    original_build = views.build_training_detail_workbook

    def track_build(*args, **kwargs):
        handle = original_build(*args, **kwargs)
        generated.append(handle)
        return handle

    monkeypatch.setattr(views, "build_training_detail_workbook", track_build)
    assert doctor.check_password("pass123456")
    client = APIClient(enforce_csrf_checks=True)
    credentials = b64encode(f"{doctor.phone}:pass123456".encode()).decode()
    client.credentials(HTTP_AUTHORIZATION=f"Basic {credentials}")
    assert "sessionid" not in client.cookies and "csrftoken" not in client.cookies
    response = download(client, project_patient)
    try:
        assert response.status_code == 403
        assert TrainingDetailExportLog.objects.count() == 0
        assert generated == []
    finally:
        response.close()


@pytest.mark.parametrize("failure", ["limit", "deadline"])
def test_business_failure_and_failed_audit_error_return_fixed_500_and_cleanup(
    session_client, export_sample, monkeypatch, caplog, failure
):
    from apps.training import export_rows as source, export_workbook as engine
    from apps.training.export_rows import ExportDeadlineError
    from zipfile import ZipFile

    paths, handles = [], []
    original_init = ExportRows.__init__
    original_temp = engine.TemporaryFile
    original_save = TrainingDetailExportLog.save

    def tracked_rows(self, *args, **kwargs):
        original_init(self, *args, **kwargs)
        paths.append(self.directory)

    def tracked_file(*args, **kwargs):
        handle = original_temp(*args, **kwargs)
        handles.append(handle)
        return handle

    def fail_failed_audit(self, *args, **kwargs):
        if self.status == "failed":
            raise OSError("SECRET_FAILED_AUDIT")
        return original_save(self, *args, **kwargs)

    monkeypatch.setattr(ExportRows, "__init__", tracked_rows)
    monkeypatch.setattr(engine, "TemporaryFile", tracked_file)
    monkeypatch.setattr(TrainingDetailExportLog, "save", fail_failed_audit)
    if failure == "limit":
        monkeypatch.setattr(source, "MAX_RECORDS", 1)
    else:

        def compression_deadline(*args, **kwargs):
            raise ExportDeadlineError("SECRET_DEADLINE")

        monkeypatch.setattr(ZipFile, "write", compression_deadline)

    response = download(session_client, export_sample.project_patient)
    assert response.status_code == 500
    assert response.data == {"detail": "导出失败，请重试"}
    assert "SECRET" not in caplog.text
    assert paths and all(not Path(path).exists() for path in paths)
    if failure == "deadline":
        assert handles
    assert all(handle.closed for handle in handles)
    assert TrainingDetailExportLog.objects.get().status == "generating"
