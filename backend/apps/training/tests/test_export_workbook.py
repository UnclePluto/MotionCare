import json
import os
import re
import time
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile

import openpyxl
import pytest
from openpyxl.worksheet._writer import ALL_TEMP_FILES

from apps.training._export_mapping import motion_row
from apps.training.export_rows import ExportDeadlineError, ExportRows
from apps.training.export_workbook import build_training_detail_workbook
from apps.training.models import TrainingRecord


@pytest.fixture
def workbook_rows():
    with ExportRows(deadline=time.monotonic() + 60) as rows:
        rows.metadata = {
            "snapshot_at": datetime(2026, 9, 9, tzinfo=UTC),
            "generated_at": datetime(2026, 9, 9, tzinfo=UTC),
            "operator": "测试医生",
        }
        rows.append(
            "训练场次",
            {
                "患者编号": 1,
                "患者姓名": "中文患者\x00",
                "训练日期": date(2026, 9, 9),
                "提交时间": datetime(2026, 9, 8, 16, tzinfo=UTC),
                "游戏得分": 0,
                "完成题数": 0,
                "项目名称": '=HYPERLINK("x")',
                "是否提前结束": False,
                "心率均值（次/分）": Decimal("70.5"),
                "备注": '=HYPERLINK("x")',
            },
        )
        rows.append(
            "训练场次", {"备注": "隐藏长文本" * 13000, "项目名称": "长文本\n" * 13000 + "\x01"}
        )
        quality = json.dumps({"doctor_note": "质" * 60000 + "\ufffe"}, ensure_ascii=False)
        rows.append(
            "运动明细", {"动作质量详情1": quality, "备注": "正常\n说明", "动作名称": quality}
        )
        rows.append("游戏逐题", {"有效作答时间（毫秒）": 0, "判定": "正确"})
        rows.append(
            "训练期间生理数据",
            {"心率（次/分）": 60, "测量时间": datetime(2026, 9, 8, 16, tzinfo=UTC)},
        )
        rows.append(
            "字段说明",
            {
                "分类": "导出元数据",
                "字段或项目": "snapshot_at",
                "说明或值": datetime(2020, 1, 1, tzinfo=UTC),
            },
        )
        yield rows


def records(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_six_sheets_safe_text_types_and_lossless_chunks(workbook_rows):
    with build_training_detail_workbook(workbook_rows, deadline=time.monotonic() + 60) as handle:
        assert handle.tell() == 0
        assert os.fstat(handle.fileno()).st_mode & 0o777 == 0o600
        book = openpyxl.load_workbook(handle, data_only=False)
        assert book.sheetnames == [
            "训练场次",
            "游戏逐题",
            "顺序选择明细",
            "运动明细",
            "训练期间生理数据",
            "字段说明",
        ]
        for sheet in book:
            assert sheet.freeze_panes == "A2"
            assert sheet.auto_filter.ref is not None
            assert not sheet.merged_cells.ranges
            assert sheet["A1"].font.bold
        sessions = records(book["训练场次"])
        # Long text expands the original column to sequential columns for every row.
        assert sessions[0]["项目名称1"] == '=HYPERLINK("x")'
        assert (
            "".join(sessions[1][f"项目名称{i}"] or "" for i in (1, 2)) == "长文本\n" * 13000 + "�"
        )
        assert sessions[0]["患者姓名"] == "中文患者�"
        assert sessions[0]["完成题数"] == 0
        header = [cell.value for cell in book["训练场次"][1]]
        assert book["训练场次"].cell(2, header.index("完成题数") + 1).number_format == "0"
        assert (
            book["训练场次"].cell(2, header.index("心率均值（次/分）") + 1).number_format == "0.0"
        )
        assert sessions[0]["是否提前结束"] is False
        assert sessions[0]["提交时间"] == datetime(2026, 9, 9)
        assert sessions[0]["训练日期"] == datetime(2026, 9, 9)
        assert sessions[0]["心率均值（次/分）"] == 70.5
        note = next(
            cell for row in book["训练场次"] for cell in row if cell.value == '=HYPERLINK("x")'
        )
        assert note.data_type == "s" and note.hyperlink is None
        motion = records(book["运动明细"])[0]
        assert motion["动作名称1"] + motion["动作名称2"] == json.dumps(
            {"doctor_note": "质" * 60000 + "�"}, ensure_ascii=False
        )
        explanations = records(book["字段说明"])
        by_key = {r["字段或项目"]: r["说明或值"] for r in explanations}
        assert by_key["数据读取截至点"] == datetime(2026, 9, 9, 8)
        assert by_key["XML非法字符替换数"] == 3
        assert by_key["字段说明"] == len(explanations) == workbook_rows.row_counts["字段说明"]
        assert by_key["导出人"] == "测试医生"
        assert "按列顺序拼接" in by_key["项目名称2"]
        assert book._external_links == [] and book.vba_archive is None


def test_empty_sheets_keep_headers_and_filters():
    with ExportRows(deadline=time.monotonic() + 60) as rows:
        with build_training_detail_workbook(rows, deadline=time.monotonic() + 60) as handle:
            book = openpyxl.load_workbook(handle)
            for sheet in list(book)[:5]:
                assert sheet.max_row == 1
                assert sheet.auto_filter.ref.endswith("1")


@pytest.mark.django_db
def test_workbook_preserves_safe_confidence_types_and_excludes_unsafe_quality_values(
    project_patient, active_prescription, prescription_action
):
    record = TrainingRecord.objects.create(
        project_patient=project_patient,
        prescription=active_prescription,
        prescription_action=prescription_action,
        training_date=date(2026, 9, 9),
        status="completed",
    )
    record.export_analysis_status = None
    cases = (
        ({"confidence_level": 0.98}, {"confidence_level": 0.98}),
        ({"confidence_level": 0}, {"confidence_level": 0}),
        ({"confidence_level": "high"}, {"confidence_level": "high"}),
        (
            {
                "confidence_level": True,
                "doctor_note": {"nested": "SECRET_NOTE"},
                "quality_flags": ["stable", 1, {"nested": "SECRET_FLAG"}],
                "object_key": "SECRET_OBJECT",
                "raw_payload": {"nested": "SECRET_RAW"},
            },
            {"quality_flags": ["stable"]},
        ),
        ({"confidence_level": float("nan")}, None),
        ({"confidence_level": float("inf")}, None),
        ({"confidence_level": {"nested": "SECRET_CONFIDENCE"}}, None),
    )

    with ExportRows(deadline=time.monotonic() + 60) as rows:
        for quality, _ in cases:
            record.motion_quality_data = quality
            rows.append("运动明细", motion_row(record, None))
        with build_training_detail_workbook(rows, deadline=time.monotonic() + 60) as handle:
            book = openpyxl.load_workbook(handle, data_only=False)
            exported = [row["动作质量详情1"] for row in rows.iter_rows("运动明细")]
            assert all("动作质量详情1" not in row for row in records(book["运动明细"]))
            assert_simplified_columns(book)

    decoded = [json.loads(value) if value is not None else None for value in exported]
    assert decoded == [expected for _, expected in cases]
    assert "SECRET_" not in json.dumps(exported, ensure_ascii=False)


@pytest.mark.parametrize("stage", ["rows", "zip", "zip_deadline"])
def test_failure_closes_xlsx_and_only_own_internal_xml(workbook_rows, monkeypatch, stage):
    from apps.training import export_workbook as engine

    handles = []
    real_temporary = engine.TemporaryFile

    def tracked(*args, **kwargs):
        handle = real_temporary(*args, **kwargs)
        handles.append(handle)
        return handle

    monkeypatch.setattr(engine, "TemporaryFile", tracked)
    real_iter = workbook_rows.iter_rows
    if stage == "rows":

        def fail_rows(name):
            yield from real_iter(name)
            if name == "运动明细":
                raise OSError("模拟行读取失败")

        monkeypatch.setattr(workbook_rows, "iter_rows", fail_rows)
    else:
        from zipfile import ZipFile

        original_write = ZipFile.write

        def fail_write(self, *args, **kwargs):
            if stage == "zip_deadline":
                monkeypatch.setattr(engine.time, "monotonic", lambda: 10**15)
                return original_write(self, *args, **kwargs)
            raise OSError("模拟压缩失败")

        monkeypatch.setattr(ZipFile, "write", fail_write)
    before = set(ALL_TEMP_FILES)
    with NamedTemporaryFile(prefix="other-export-", delete=False) as other:
        try:
            with pytest.raises((OSError, ExportDeadlineError)):
                build_training_detail_workbook(workbook_rows, deadline=time.monotonic() + 60)
            assert handles and all(handle.closed for handle in handles)
            assert set(ALL_TEMP_FILES) == before
            assert Path(other.name).exists()
        finally:
            os.unlink(other.name)


def test_writer_constructor_failure_does_not_leave_internal_xml(workbook_rows, monkeypatch):
    from openpyxl.worksheet import _writer

    before = set(ALL_TEMP_FILES)

    def broken_stream(self):
        raise OSError("模拟XML初始化失败")
        yield

    monkeypatch.setattr(_writer.WorksheetWriter, "get_stream", broken_stream)
    try:
        with pytest.raises(OSError):
            build_training_detail_workbook(workbook_rows, deadline=time.monotonic() + 60)
        assert set(ALL_TEMP_FILES) == before
    finally:
        # Red-run cleanup only: these are exactly this test's leaked writer paths.
        for path in set(ALL_TEMP_FILES) - before:
            Path(path).unlink(missing_ok=True)
            ALL_TEMP_FILES.remove(path)


@pytest.mark.parametrize("stage", ["success", "writer_init", "zip"])
def test_private_xml_directory_files_are_private_and_cleaned(workbook_rows, monkeypatch, stage):
    from apps.training import export_workbook as engine

    directories = []
    original_temp = engine.TemporaryDirectory
    original_init = engine._OwnedWorksheetWriter.__init__

    def directory(*args, **kwargs):
        temp = original_temp(*args, **kwargs)
        directories.append(temp.name)
        assert os.stat(temp.name).st_mode & 0o777 == 0o700
        return temp

    def writer_init(self, sheet, out):
        assert os.stat(out).st_mode & 0o777 == 0o600
        if stage == "writer_init":
            raise OSError("模拟初始化中断")
        return original_init(self, sheet, out=out)

    monkeypatch.setattr(engine, "TemporaryDirectory", directory)
    monkeypatch.setattr(engine._OwnedWorksheetWriter, "__init__", writer_init)
    if stage == "zip":
        from zipfile import ZipFile

        def fail(*args, **kwargs):
            raise OSError("模拟压缩中断")

        monkeypatch.setattr(ZipFile, "write", fail)
    if stage == "success":
        with build_training_detail_workbook(workbook_rows, deadline=time.monotonic() + 60):
            pass
    else:
        with pytest.raises(OSError):
            build_training_detail_workbook(workbook_rows, deadline=time.monotonic() + 60)
    assert directories and all(not Path(path).exists() for path in directories)


OMITTED_LABELS = {
    "起止时间可用性",
    "游戏得分",
    "逐题完整性",
    "备注",
    "窗口口径",
    "心率数据可用性",
    "血压数据可用性",
    "血氧数据可用性",
    "数据口径说明",
    "动作质量详情1",
}


def omitted_label(label):
    return label in OMITTED_LABELS or re.fullmatch(r"(?:备注|动作质量详情)\d*", label) is not None


def assert_simplified_columns(book):
    for sheet in book:
        assert not any(omitted_label(cell.value) for cell in sheet[1])
    definitions = [
        row["字段或项目"] for row in records(book["字段说明"]) if row["分类"] == "字段定义"
    ]
    assert not any(omitted_label(label) for label in definitions)


def test_all_sheets_and_definitions_omit_requested_columns(workbook_rows):
    with build_training_detail_workbook(workbook_rows, deadline=time.monotonic() + 60) as handle:
        book = openpyxl.load_workbook(handle)
        assert_simplified_columns(book)
        headers = {cell.value for sheet in book for cell in sheet[1]}
        assert {
            "有效作答时间（毫秒）",
            "有效选择耗时（毫秒）",
            "实际交换次数",
            "拼图点击次数",
            "窗口开始",
            "窗口结束",
        } <= headers
        assert len(list(workbook_rows.iter_rows("运动明细"))[0]["动作质量详情1"]) == 32767
