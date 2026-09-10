"""合成数据的患者写入→医生 Session 导出验收；不访问患者业务库。"""

from datetime import UTC, date, datetime, timedelta
from io import BytesIO
import json
import os
from pathlib import Path
import time
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZipFile

import openpyxl
import pytest
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from rest_framework.test import APIClient

from apps.patient_app.services import bind_project_patient_with_code, create_binding_code
from apps.patients.models import Patient
from apps.prescriptions.models import ActionLibraryItem, Prescription
from apps.studies.models import ProjectPatient, StudyProject
from apps.training.export_rows import ExportLimitError, ExportRows, prepare_export_rows
from apps.training.export_schema import SHEET_HEADERS
from apps.training.tests.test_export_workbook import assert_simplified_columns, omitted_label
from apps.training.export_scope import ExportFilter
from apps.training.models import (
    MotionAnalysisJob,
    TrainingDetailExportLog,
    TrainingRecord,
    TrainingVideo,
)
from apps.wearables.models import WearableDevice, WearableMeasurement

pytestmark = pytest.mark.django_db(transaction=True)
GAME_CODES = (
    "game-memory-color-sequence",
    "game-memory-pattern-sequence",
    "game-executive-inhibition",
    "game-executive-category-switch",
    "game-audiovisual-sound-discrimination",
    "game-audiovisual-puzzle",
)
GAME_NAMES = ("颜色顺序", "图案顺序", "反应抑制", "分类转换", "声音辨别", "图片拼图")
DAY = date(2026, 9, 9)


def patient_client(pp, doctor):
    code, _ = create_binding_code(project_patient=pp, created_by=doctor)
    token, _ = bind_project_patient_with_code(code, wx_openid=f"acceptance-{pp.pk}")
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")
    return client


def doctor_client(doctor):
    client = APIClient(enforce_csrf_checks=True)
    client.get("/api/auth/csrf/")
    result = client.post(
        "/api/auth/login/",
        {"phone": doctor.phone, "password": "pass123456"},
        format="json",
        HTTP_X_CSRFTOKEN=client.cookies["csrftoken"].value,
    )
    assert result.status_code == 204
    client.credentials(HTTP_X_CSRFTOKEN=client.cookies["csrftoken"].value)
    return client


def export(client, pp, **filters):
    return client.post(
        f"/api/training/tracking/patients/{pp.patient_id}/export/",
        {"project_patient": pp.pk, "range": "all", **filters},
        format="json",
    )


def workbook(response):
    assert response.status_code == 200, getattr(response, "data", None)
    content = b"".join(response.streaming_content)
    response.close()
    return openpyxl.load_workbook(BytesIO(content)), content


def records(sheet):
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def submit(client, action, difficulty="简单", questions=None, **overrides):
    code = action.action_library_item.source_key
    if questions is None:
        questions = [
            {
                "question_index": i,
                "game_code": code,
                "difficulty": difficulty,
                "response_duration_ms": ms,
                "is_correct": i == 1 or code == GAME_CODES[-1],
                "result_type": "answered" if i == 1 or code == GAME_CODES[-1] else "timeout",
                "swap_count": i - 1 if code == GAME_CODES[-1] else None,
            }
            for i, ms in ((1, 0), (2, 1200))
        ]
    payload = {
        "prescription_action": action.pk,
        "training_date": DAY.isoformat(),
        "status": "completed",
        "actual_duration_minutes": 1,
        "score": "0.00",
        "note": "=合成验收备注",
        "client_session_id": str(uuid4()),
        "form_data": {
            "difficulty": difficulty,
            "raw_detail": {
                "session_duration_seconds": 10,
                "ended_early": False,
                "ended_by": "timer",
                "rounds": [{"round_index": 90, "response_ms": 999, "correct": True}],
            },
        },
        "question_results": questions,
        **overrides,
    }
    response = client.post("/api/patient-app/training-records/", payload, format="json")
    assert response.status_code == 201, response.data
    return TrainingRecord.objects.get(pk=response.data["id"]), payload


@pytest.fixture
def acceptance_sample(project_patient, active_prescription, doctor):
    pp = project_patient
    pp.patient.name = "合成患者·导出验收"
    pp.patient.save(update_fields=["name"])
    pp.project.name = "合成康复研究项目"
    pp.project.save(update_fields=["name"])
    client = patient_client(pp, doctor)
    submitted, actions = [], []
    for index, (code, name) in enumerate(zip(GAME_CODES, GAME_NAMES)):
        item, _ = ActionLibraryItem.objects.get_or_create(
            source_key=code,
            defaults={
                "name": name,
                "internal_type": "game",
                "training_type": "认知训练",
                "action_type": "训练",
            },
        )
        action = active_prescription.add_action_snapshot(
            item, difficulty="简单", duration_minutes=1
        )
        actions.append(action)
        submitted.append(submit(client, action, ("简单", "中等", "困难")[index % 3]))
    empty, _ = submit(client, actions[0], questions=[])
    legacy_payload = {
        "prescription_action": actions[0].pk,
        "training_date": "2026-09-08",
        "status": "completed",
        "actual_duration_minutes": 2,
        "form_data": {
            "difficulty": "简单",
            "accuracy_rate": 75,
            "raw_detail": {
                "rounds": [
                    {"round_index": 1, "response_ms": 0, "correct": False},
                    {"round_index": 3, "response_ms": 2400, "correct": True},
                    {"round_index": 3, "response_ms": 1, "correct": True},
                    {"round_index": 4, "response_ms": "bad", "correct": True},
                ]
            },
        },
    }
    response = client.post("/api/patient-app/training-records/", legacy_payload, format="json")
    assert response.status_code == 201, response.data
    legacy = TrainingRecord.objects.get(pk=response.data["id"])
    active_prescription.status = Prescription.Status.ARCHIVED
    active_prescription.save(update_fields=["status"])
    current = Prescription.objects.create(
        project_patient=pp,
        version=2,
        opened_by=doctor,
        status="active",
        effective_at=timezone.now(),
    )
    # 另一个项目也走患者提交链路，确保单项目导出不会混入。
    other_project = StudyProject.objects.create(name="合成另一项目", created_by=doctor)
    other_pp = ProjectPatient.objects.create(project=other_project, patient=pp.patient)
    other_rx = Prescription.objects.create(
        project_patient=other_pp,
        version=1,
        opened_by=doctor,
        status="active",
        effective_at=timezone.now(),
    )
    other_action = other_rx.add_action_snapshot(actions[0].action_library_item, difficulty="简单")
    other, _ = submit(patient_client(other_pp, doctor), other_action)
    started = datetime(2026, 9, 8, 15, 59, 50, tzinfo=UTC)
    long_note = "=医生说明：合成质量😀" * 4000
    motions, videos = [], []
    for index, (key, name, ai, source) in enumerate(
        (
            ("motion-resistance-shoulder-press", "肩部推举", True, "algorithm"),
            ("motion-resistance-shoulder-press", "肩部推举", True, "doctor"),
            ("motion-aerobic-high-knee", "椰林步道模拟", False, "doctor"),
        )
    ):
        item, _ = ActionLibraryItem.objects.get_or_create(
            source_key=key,
            defaults={
                "name": name,
                "internal_type": "motion",
                "training_type": "运动训练",
                "action_type": "训练",
                "has_ai_supervision": ai,
            },
        )
        # 合成夹具显式设置能力；迁移初始目录的肩推值不代表运行时目录同步结果。
        item.has_ai_supervision = ai
        item.save(update_fields=["has_ai_supervision"])
        action = current.add_action_snapshot(item, duration_minutes=3)
        record = TrainingRecord.objects.create(
            project_patient=pp,
            prescription=current,
            prescription_action=action,
            training_date=DAY,
            status="completed",
            motion_total_count=5,
            motion_standard_count=4,
            motion_nonstandard_count=1,
            motion_result_source=source,
            motion_result_updated_at=started,
            motion_result_updated_by=doctor if source == "doctor" else None,
            motion_quality_data={
                "doctor_note": long_note if index == 1 else "合成说明",
                "object_key": "SECRET_OBJECT",
                "raw_payload": {"secret": "SECRET_RAW"},
            },
        )
        video = TrainingVideo.objects.create(
            project_patient=pp,
            prescription=current,
            prescription_action=action,
            training_record=record,
            training_date=DAY,
            training_started_at=started,
            expected_duration_seconds=180,
            actual_duration_seconds=0,
            training_ended_at=None if index == 1 else started + timedelta(seconds=20),
            duration_seconds=77,
            status="attached",
            object_key=f"SECRET_VIDEO/{uuid4()}",
        )
        if ai:
            MotionAnalysisJob.objects.create(
                training_video=video,
                training_record=record,
                prescription_action=action,
                project_patient=pp,
                status="succeeded" if index == 0 else "failed",
                failure_reason="SECRET_STACK",
            )
        motions.append(record)
        videos.append(video)
    device = WearableDevice.objects.create(
        provider="miwitracker",
        external_device_id="synthetic-export",
        identifier_type="device_id",
        short_code="9999",
    )
    outsider = Patient.objects.create(name="合成其他患者", primary_doctor=doctor)

    def measure(at, metric="heart_rate", patient=None, attribution="attributed", **values):
        return WearableMeasurement.objects.create(
            provider="miwitracker",
            patient=patient or pp.patient,
            device=device,
            metric_type=metric,
            measured_at=at,
            attribution_status=attribution,
            source_fingerprint=str(uuid4()),
            raw_payload={"secret": "SECRET_DEVICE"},
            **values,
        )

    accepted = [
        measure(started, heart_rate=60),
        measure(started + timedelta(seconds=480), heart_rate=81),
        measure(started, "blood_pressure", systolic=120, diastolic=80),
        measure(started, "blood_oxygen", blood_oxygen=0),
    ]
    excluded = [
        measure(started - timedelta(microseconds=1), heart_rate=199),
        measure(started + timedelta(seconds=480, microseconds=1), heart_rate=199),
        measure(started, patient=outsider, heart_rate=199),
        measure(started, attribution="ambiguous", heart_rate=199),
        measure(started, "blood_pressure", systolic=120),
    ]
    return SimpleNamespace(
        pp=pp,
        doctor=doctor,
        client=doctor_client(doctor),
        submitted=submitted,
        empty=empty,
        legacy=legacy,
        other=other,
        motions=motions,
        videos=videos,
        started=started,
        accepted=accepted,
        excluded=excluded,
        long_note=long_note,
    )


def test_patient_api_to_session_export_known_values_and_actual_sample(acceptance_sample):
    s = acceptance_sample
    StudyProject.objects.filter(pk=s.pp.project_id).update(status="completed")
    response = export(s.client, s.pp)
    assert response["Cache-Control"] == "no-store"
    book, content = workbook(response)
    assert book.sheetnames == list(SHEET_HEADERS)
    assert_simplified_columns(book)
    sheets = {name: records(book[name]) for name in book.sheetnames}
    sessions = {row["训练记录编号"]: row for row in sheets["训练场次"]}
    expected_ids = (
        {record.pk for record, _ in s.submitted}
        | {s.empty.pk, s.legacy.pk}
        | {r.pk for r in s.motions}
    )
    assert set(sessions) == expected_ids and s.other.pk not in sessions
    assert {row["处方版本"] for row in sessions.values()} == {1, 2}
    questions = sheets["游戏逐题"]
    expected_active_count = sum(len(payload["question_results"]) for _, payload in s.submitted)
    assert {row["游戏编码"] for row in questions} == set(GAME_CODES)
    assert (
        sum(row["采集版本"] == "active_response_v1" for row in questions) == expected_active_count
    )
    assert len(questions) == expected_active_count
    assert all("历史作答时间（毫秒）" not in row for row in questions)
    assert not any(row["训练记录编号"] == s.empty.pk or row["题号"] == 90 for row in questions)
    for record, payload in s.submitted:
        rows = [r for r in questions if r["训练记录编号"] == record.pk]
        assert [(r["题号"], r["有效作答时间（毫秒）"], r["实际交换次数"]) for r in rows] == [
            (q["question_index"], q["response_duration_ms"], q["swap_count"])
            for q in payload["question_results"]
        ]
        assert {r["实际难度"] for r in rows} == {payload["form_data"]["difficulty"]}
        summary = sessions[record.pk]
        assert record.note == "=合成验收备注" and record.score == 0
        assert summary["完成题数"] == 2 and summary["是否提前结束"] is False
        assert summary["实际训练时长（秒）"] == 10
        assert summary["训练开始时间"] is None and summary["心率均值（次/分）"] is None
        assert summary["正确率（%）"] == (None if rows[0]["游戏编码"] == GAME_CODES[-1] else 50)
    legacy = [r for r in questions if r["训练记录编号"] == s.legacy.pk]
    assert legacy == []
    assert sessions[s.legacy.pk]["逐题记录数"] == 0
    assert sessions[s.legacy.pk]["规范计时题数"] == 0
    assert sessions[s.legacy.pk]["正确率（%）"] == 75
    assert "历史计时题数" not in sessions[s.legacy.pk]
    assert s.legacy.question_results.count() == 2
    assert sessions[s.legacy.pk]["实际训练时长（秒）"] == 120
    physiology = sheets["训练期间生理数据"]
    assert len(physiology) == len(s.accepted) * len(s.motions)
    assert {r["测量记录编号"] for r in physiology} == {m.pk for m in s.accepted}
    for record in s.motions:
        summary = sessions[record.pk]
        assert summary["窗口开始"] == datetime(2026, 9, 8, 23, 59, 50)
        assert summary["窗口结束"] == datetime(2026, 9, 9, 0, 7, 50)
        assert (
            summary["心率读数数量"],
            summary["心率均值（次/分）"],
            summary["血压配对数量"],
            summary["血氧均值（%）"],
        ) == (2, 70.5, 1, 0)
    motion = {r["训练记录编号"]: r for r in sheets["运动明细"]}
    assert [motion[r.pk]["结果来源"] for r in s.motions] == ["算法", "医生", "医生"]
    assert motion[s.motions[2].pk]["分析状态"] == "不支持AI"
    assert motion[s.motions[1].pk]["分析状态"] == "分析失败"
    assert motion[s.motions[0].pk]["实际训练秒数"] == 0
    assert motion[s.motions[0].pk]["视频文件秒数"] == 77
    quality_keys = [c.value for c in book["运动明细"][1] if c.value.startswith("动作质量详情")]
    assert quality_keys == []
    s.motions[1].refresh_from_db()
    assert s.motions[1].motion_quality_data["doctor_note"] == s.long_note
    assert_simplified_columns(book)
    # 遍历每张宽表的全部列和全部单元格，而非只检查截图可见区域。
    for sheet in book:
        headers = [c.value for c in sheet[1]]
        assert len(headers) == len(set(headers))
        assert {label for label in SHEET_HEADERS[sheet.title] if not omitted_label(label)} <= set(
            headers
        )
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == sheet.dimensions
        assert not sheet.merged_cells.ranges
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                assert cell.data_type != "f" and cell.hyperlink is None
                if isinstance(cell.value, str):
                    assert cell.data_type == "s" and len(cell.value) <= 32767
                    assert "SECRET_" not in cell.value
                elif isinstance(cell.value, datetime):
                    assert cell.data_type == "d" and cell.number_format in (
                        "yyyy-mm-dd",
                        "yyyy-mm-dd hh:mm:ss",
                    )
                elif isinstance(cell.value, bool):
                    assert cell.data_type == "b"
                elif isinstance(cell.value, (int, float)):
                    assert cell.data_type == "n"
    with ZipFile(BytesIO(content)) as archive:
        assert not any(
            "externalLinks" in name or "vbaProject" in name for name in archive.namelist()
        )
    audit = TrainingDetailExportLog.objects.get(status="succeeded")
    assert audit.row_counts == {name: len(rows) for name, rows in sheets.items()}
    if os.environ.get("MOTIONCARE_EXPORT_SAMPLE"):
        destination = Path(os.environ["MOTIONCARE_EXPORT_SAMPLE"])
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)
        destination.with_suffix(".verification.json").write_text(
            json.dumps(
                {
                    "source": "合成数据：患者Bearer API写入→医生Session+CSRF导出API",
                    "rows": audit.row_counts,
                    "columns": {s.title: s.max_column for s in book},
                    "active_question_count": expected_active_count,
                    "quality_segments": len(quality_keys),
                    "excluded_legacy_question_count": 2,
                    "only_valid_active_response_durations": True,
                    "all_columns_checked": True,
                    "formula_and_links": 0,
                    "requested_columns_omitted_from_all_sheets_and_definitions": True,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    book.close()


def test_date_endpoints_same_day_multiple_sessions_and_empty_six_sheets(acceptance_sample):
    s = acceptance_sample
    book, _ = workbook(
        export(s.client, s.pp, range="custom", start_date="2026-09-09", end_date="2026-09-09")
    )
    assert len(records(book["训练场次"])) == 10
    assert {r["训练日期"] for r in records(book["训练场次"])} == {datetime(2026, 9, 9)}
    book.close()
    book, _ = workbook(
        export(s.client, s.pp, range="custom", start_date="2026-09-08", end_date="2026-09-09")
    )
    assert len(records(book["训练场次"])) == 11
    book.close()
    book, _ = workbook(
        export(s.client, s.pp, range="custom", start_date="2026-09-07", end_date="2026-09-07")
    )
    assert book.sheetnames == list(SHEET_HEADERS)
    assert [book[name].max_row for name in book.sheetnames[:5]] == [1] * 5
    assert book["字段说明"].max_row > 1
    book.close()


def test_real_ten_thousand_record_capacity_and_overflow(
    project_patient, prescription_action, doctor
):
    started = time.monotonic()
    TrainingRecord.objects.bulk_create(
        [
            TrainingRecord(
                project_patient=project_patient,
                prescription=prescription_action.prescription,
                prescription_action=prescription_action,
                training_date=DAY,
                status="completed",
            )
            for _ in range(10000)
        ]
    )
    filters = ExportFilter(project_patient.pk, "all", None, None)
    with (
        CaptureQueriesContext(connection) as queries,
        prepare_export_rows(
            project_patient=project_patient, filters=filters, deadline=time.monotonic() + 60
        ) as rows,
    ):
        assert rows.row_counts["训练场次"] == 10000
        assert rows.row_counts["运动明细"] == 10000
        assert sum(1 for _ in rows.iter_rows("训练场次")) == 10000
        query_count = len(queries)
    elapsed = time.monotonic() - started
    assert elapsed < 60 and 1 <= query_count <= 610
    TrainingRecord.objects.create(
        project_patient=project_patient,
        prescription=prescription_action.prescription,
        prescription_action=prescription_action,
        training_date=DAY,
        status="completed",
    )
    response = export(doctor_client(doctor), project_patient)
    assert response.status_code == 400 and "application/json" in response["Content-Type"]
    assert "缩小" in response.data["detail"]
    assert TrainingDetailExportLog.objects.get().status == "failed"
    print(f"容量实测：10000场、20000数据行，{query_count}次SQL，{elapsed:.2f}秒；10001场拒绝")


def test_real_sheet_and_total_capacity_without_relaxing_limits():
    started = time.monotonic()
    with ExportRows(deadline=started + 60) as rows:
        directory = rows.directory
        for _ in range(200000):
            rows.append("游戏逐题", {"题号": 1})
        with pytest.raises(ExportLimitError):
            rows.append("游戏逐题", {"题号": 1})
        for _ in range(200000):
            rows.append("训练期间生理数据", {"心率（次/分）": 60})
        for _ in range(100000):
            rows.append("运动明细", {"动作总次数": 0})
        assert sum(rows.row_counts.values()) == 500000
        with pytest.raises(ExportLimitError):
            rows.append("运动明细", {"动作总次数": 0})
    assert not Path(directory).exists()
    print(f"真实行上限：单表200000、总500000，{time.monotonic() - started:.2f}秒，溢出拒绝并清理")


def test_upgrade_from_pre_feature_leaf_preserves_legacy_json(
    project_patient, active_prescription, capsys
):
    executor = MigrationExecutor(connection)
    latest = executor.loader.graph.leaf_nodes()
    # 由迁移图定位本功能前实际叶节点，避免误把0016当成旧版。
    before = [
        (
            "training",
            next(
                name
                for app, name in executor.loader.graph.nodes
                if app == "training" and name.startswith("0015_")
            ),
        )
    ]
    try:
        executor.migrate(before)
        old_apps = executor.loader.project_state(before).apps
        Item = old_apps.get_model("prescriptions", "ActionLibraryItem")
        Action = old_apps.get_model("prescriptions", "PrescriptionAction")
        Record = old_apps.get_model("training", "TrainingRecord")
        item = Item.objects.create(
            source_key=GAME_CODES[0],
            name="合成颜色顺序",
            internal_type="game",
            training_type="认知训练",
        )
        action = Action.objects.create(
            prescription_id=active_prescription.pk,
            action_library_item_id=item.pk,
            action_name_snapshot="合成颜色顺序",
            training_type_snapshot="认知训练",
            internal_type_snapshot="game",
            difficulty="简单",
        )
        source = {
            "difficulty": "中等",
            "raw_detail": {
                "rounds": [
                    {"round_index": 1, "response_ms": 0, "correct": False},
                    {"round_index": 3, "response_ms": 2400, "correct": True},
                    {"round_index": 3, "response_ms": 1, "correct": True},
                    {"round_index": 4, "response_ms": "bad", "correct": True},
                ]
            },
        }
        record = Record.objects.create(
            project_patient_id=project_patient.pk,
            prescription_id=active_prescription.pk,
            prescription_action_id=action.pk,
            training_date=DAY,
            status="completed",
            form_data=source,
        )
        capsys.readouterr()
        executor = MigrationExecutor(connection)
        executor.migrate(latest)
        output = capsys.readouterr().out
        assert "scanned=1 created=2 skipped=2" in output
        record = TrainingRecord.objects.get(pk=record.pk)
        assert record.form_data == source
        assert list(
            record.question_results.order_by("question_index").values_list(
                "question_index", "response_duration_ms", "capture_version"
            )
        ) == [(1, 0, "legacy_wall_clock_v0"), (3, 2400, "legacy_wall_clock_v0")]
        print(f"迁移演练 {before[0][1]}→0018：scanned=1 created=2 skipped=2；来源JSON逐值未变")
    finally:
        MigrationExecutor(connection).migrate(latest)


def test_v2_selection_details_real_api_and_sample(acceptance_sample):
    s = acceptance_sample
    rx = Prescription.objects.get(project_patient=s.pp, status="active")
    patient = patient_client(s.pp, s.doctor)
    submitted = []
    for code in GAME_CODES:
        action = rx.add_action_snapshot(
            ActionLibraryItem.objects.get(source_key=code), difficulty="简单", duration_minutes=1
        )

        def question(index, result="answered"):
            sequence = code in GAME_CODES[:2]
            tokens = (
                ["blue", "green", "yellow"] if code == GAME_CODES[0] else ["sun", "coconut", "boat"]
            )
            steps = (
                [
                    {
                        "step_index": i + 1,
                        "selected_value": token if i != 1 else tokens[0],
                        "expected_value": token,
                        "response_duration_ms": 100 * (i + 1),
                        "is_correct": i != 1,
                    }
                    for i, token in enumerate(tokens)
                ]
                if sequence
                else []
            )
            if result != "answered":
                steps = steps[:1]
            return {
                "question_index": index,
                "game_code": code,
                "difficulty": "简单",
                "response_duration_ms": 900,
                "is_correct": not sequence and result == "answered",
                "result_type": result,
                "swap_count": 2 if code == GAME_CODES[-1] else None,
                "capture_version": "active_response_v2",
                "expected_step_count": 3 if sequence else None,
                "click_count": 6 if code == GAME_CODES[-1] else None,
                "selection_steps": steps,
            }

        questions = [question(1)]
        if code in GAME_CODES[:2]:
            questions += [question(2, "timeout"), question(3, "interrupted")]
        record, payload = submit(patient, action, questions=questions)
        submitted.append((record, payload))
    book, content = workbook(export(s.client, s.pp))
    assert book.sheetnames == [
        "训练场次",
        "游戏逐题",
        "顺序选择明细",
        "运动明细",
        "训练期间生理数据",
        "字段说明",
    ]
    assert_simplified_columns(book)
    sheets = {name: records(book[name]) for name in book.sheetnames}
    sessions = {r["训练记录编号"]: r for r in sheets["训练场次"]}
    steps = sheets["顺序选择明细"]
    assert len(steps) == 10
    assert {r["所选内容"] for r in steps} == {"蓝色", "黄色", "太阳", "小船"}
    for record, payload in submitted:
        qs = [q for q in sheets["游戏逐题"] if q["训练记录编号"] == record.pk]
        sequence = payload["question_results"][0]["game_code"] in GAME_CODES[:2]
        assert len(qs) == (3 if sequence else 1)
        assert sessions[record.pk]["完成题数"] == (2 if sequence else 1)
        assert sessions[record.pk]["错误题数"] == (2 if sequence else 0)
        if sequence:
            assert [q["题目状态"] for q in qs] == ["已完成", "未完成", "未完成"]
            assert [q["已选择步数"] for q in qs] == [3, 1, 1]
            assert qs[-1]["判定"] == "未完成"
            assert sessions[record.pk]["正确率（%）"] == 0
            own_steps = [r for r in steps if r["训练记录编号"] == record.pk]
            assert [
                (r["题号"], r["选择序号"], r["有效选择耗时（毫秒）"], r["判定"]) for r in own_steps
            ] == [
                (1, 1, 100, "正确"),
                (1, 2, 200, "错误"),
                (1, 3, 300, "正确"),
                (2, 1, 100, "正确"),
                (3, 1, 100, "正确"),
            ]
        if payload["question_results"][0]["game_code"] == GAME_CODES[-1]:
            assert qs[0]["拼图点击次数"] == 6 and qs[0]["实际交换次数"] == 2
    v1 = [q for q in sheets["游戏逐题"] if q["采集版本"] == "active_response_v1"]
    assert len(v1) == 12 and all(q["已选择步数"] is None and q["拼图点击次数"] is None for q in v1)
    assert all(r["有效作答时间（毫秒）"] in (0, 1200) for r in v1)
    audit = TrainingDetailExportLog.objects.get(status="succeeded")
    assert audit.format_version == "training_detail_v2"
    assert audit.row_counts == {name: len(rows) for name, rows in sheets.items()}
    explanations = {
        r["字段或项目"]: r["说明或值"] for r in sheets["字段说明"] if r["分类"] == "各表行数"
    }
    assert explanations == audit.row_counts
    for sheet in book:
        assert all(
            cell.data_type != "f" and cell.hyperlink is None for row in sheet for cell in row
        )
    destination_value = os.environ.get("MOTIONCARE_SELECTION_SAMPLE")
    if destination_value:
        destination = Path(destination_value)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(content)
        destination.with_suffix(".verification.json").write_text(
            json.dumps(
                {
                    "source": "合成数据：患者Bearer API写入→医生Session+CSRF导出API",
                    "format_version": audit.format_version,
                    "rows": audit.row_counts,
                    "columns": {sheet.title: sheet.max_column for sheet in book},
                    "selection_step_count": len(steps),
                    "completed_wrong_sequence": True,
                    "interrupted_and_timeout_retained": True,
                    "single_choice_and_puzzle_clicks_verified": True,
                    "legacy_filtered_v1_duration_preserved": True,
                    "formula_and_links": 0,
                    "requested_columns_omitted_from_all_sheets_and_definitions": True,
                },
                ensure_ascii=False,
                indent=2,
            )
        )
    book.close()
