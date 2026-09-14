"""六表流式 Excel 编码；仅持有本次请求的临时资源。"""

from contextlib import suppress
from datetime import date, datetime
from decimal import Decimal
import os
import re
from pathlib import Path
from tempfile import TemporaryDirectory, TemporaryFile
import time
from typing import BinaryIO
from zoneinfo import ZoneInfo
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter
from openpyxl.worksheet._writer import WorksheetWriter

from .export_rows import ExportDeadlineError, ExportRows, TEXT_CHUNK_SIZE
from .export_schema import FIELD_DEFINITIONS, OMITTED_EXPORT_COLUMNS, SHEET_HEADERS

SHANGHAI = ZoneInfo("Asia/Shanghai")
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
WRAPPED = Alignment(vertical="top", wrap_text=True)
METADATA_LABELS = {
    "format_version": "格式版本",
    "timezone": "时区",
    "snapshot_at": "数据读取截至点",
    "generated_at": "生成时间",
    "operator": "导出人",
    "operator_id": "导出人编号",
    "patient_id": "患者编号",
    "project_id": "项目编号",
    "project_patient_id": "项目患者编号",
    "range": "请求日期范围",
    "requested_start_date": "请求开始日期",
    "requested_end_date": "请求结束日期",
    "actual_start_date": "实际开始日期",
    "actual_end_date": "实际结束日期",
    "group_scope": "分组口径",
    "xml_replacement_count": "XML非法字符替换数",
}
HEADER_FILL = PatternFill("solid", fgColor="E8EEF5")


def _check_deadline(deadline):
    if time.monotonic() >= deadline:
        raise ExportDeadlineError("导出生成超时，请缩小日期范围")


class _DeadlineFile:
    def __init__(self, file, deadline):
        self.file, self.deadline = file, deadline

    def write(self, data):
        _check_deadline(self.deadline)
        return self.file.write(data)

    def seek(self, *args):
        return self.file.seek(*args)

    def tell(self):
        return self.file.tell()

    def flush(self):
        return self.file.flush()


class _Encoder:
    def __init__(self):
        self.replacements = 0

    def cell(self, sheet, value, field="", *, header=False):
        if isinstance(value, str):
            value, count = INVALID_XML.subn("�", value)
            self.replacements += count
        if isinstance(value, datetime):
            value = value.astimezone(SHANGHAI).replace(tzinfo=None)
        if isinstance(value, Decimal):
            value = float(value)
        cell = WriteOnlyCell(sheet, value=value)
        cell.alignment = WRAPPED
        if isinstance(value, str):
            cell.data_type = "s"
            cell.hyperlink = None
        elif isinstance(value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
        elif isinstance(value, date):
            cell.number_format = "yyyy-mm-dd"
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            cell.number_format = (
                "0.0"
                if "均值" in field or field == "正确率（%）"
                else ("0" if isinstance(value, int) else "General")
            )
        if header:
            cell.font = Font(bold=True, color="24364B")
            cell.fill = HEADER_FILL
        return cell


def _columns(rows, name):
    """(输出名, 原字段, 分段序号, 总段数)，质量字段已由 ExportRows 切分。"""
    columns = []
    for base in SHEET_HEADERS[name]:
        if base in OMITTED_EXPORT_COLUMNS:
            continue
        count = rows.text_chunk_counts[name].get(base, 1)
        quality = name == "运动明细" and base == "动作质量详情1"
        if quality:
            count = rows.quality_chunk_count
        for index in range(count):
            title = (
                f"动作质量详情{index + 1}"
                if quality
                else (f"{base}{index + 1}" if count > 1 else base)
            )
            columns.append((title, title if quality else base, index if not quality else 0, count))
    return columns


def _width(name):
    if any(word in name for word in ("备注", "说明", "详情", "口径", "可用性")):
        return 44
    if any(word in name for word in ("日期", "时间", "开始", "结束", "窗口")):
        return 20
    if any(word in name for word in ("名称", "姓名", "编码", "字段或项目")):
        return 24
    return 16


def _explanations(rows, columns):
    explanations = []
    metadata_seen = set()
    for row in rows.iter_rows("字段说明"):
        row = dict(row)
        category, key = row["分类"], row["字段或项目"]
        if category == "各表行数":
            continue
        if category == "字段定义" and key in OMITTED_EXPORT_COLUMNS:
            continue
        if category == "导出元数据" and key in rows.metadata:
            row["说明或值"] = rows.metadata[key]
            metadata_seen.add(key)
        explanations.append(row)
    for key, value in rows.metadata.items():
        if key not in metadata_seen and key != "row_counts":
            explanations.append({"分类": "导出元数据", "字段或项目": key, "说明或值": value})
    for sheet_columns in columns.values():
        for title, base, index, count in sheet_columns:
            if count > 1:
                definition = FIELD_DEFINITIONS.get(base, FIELD_DEFINITIONS["动作质量详情1"])
                explanations.append(
                    {
                        "分类": "字段定义",
                        "字段或项目": title,
                        "说明或值": f"{definition} 共{count}段，按列顺序拼接还原。",
                    }
                )
    for name in SHEET_HEADERS:
        explanations.append(
            {"分类": "各表行数", "字段或项目": name, "说明或值": rows.row_counts[name]}
        )
    explanations.append({"分类": "数据处理", "字段或项目": "xml_replacement_count", "说明或值": 0})
    rows.row_counts["字段说明"] = len(explanations)
    rows.metadata["row_counts"] = dict(rows.row_counts)
    for row in explanations:
        if row["分类"] == "各表行数" and row["字段或项目"] == "字段说明":
            row["说明或值"] = len(explanations)
    return explanations


class _OwnedWorksheetWriter(WorksheetWriter):
    def cleanup(self):
        # Files live in our private directory; openpyxl's process-global registry
        # never owns these paths, including failures inside writer construction.
        Path(self.out).unlink(missing_ok=True)


def _cleanup_workbook(book):
    # openpyxl 3.x does not remove worksheet XML when append/save is interrupted.
    # Only touch writers owned by this workbook; never sweep its process-global registry.
    for sheet in book:
        writer = sheet._writer
        if writer is None:
            continue
        if sheet._rows is not None:
            with suppress(Exception):
                sheet._rows.close()
        with suppress(Exception):
            writer.close()
        if os.path.exists(writer.out):
            writer.cleanup()
    book.close()


def build_training_detail_workbook(rows: ExportRows, *, deadline: float) -> BinaryIO:
    """成功返回已归零且权限0600的临时文件；失败清理本次所有工作簿资源。"""
    handle = TemporaryFile(mode="w+b")
    book = None
    xml_temp = None
    encoder = _Encoder()
    try:
        book = Workbook(write_only=True)
        xml_temp = TemporaryDirectory(prefix="motioncare-training-xlsx-")
        _check_deadline(deadline)
        columns = {name: _columns(rows, name) for name in SHEET_HEADERS}
        explanations = _explanations(rows, columns)
        for name, sheet_columns in columns.items():
            sheet = book.create_sheet(name)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(sheet_columns))}{max(1, rows.row_counts[name] + 1)}"
            )
            for position, (title, _, _, _) in enumerate(sheet_columns, 1):
                sheet.column_dimensions[get_column_letter(position)].width = _width(title)
            sheet.row_dimensions[1].height = 42
            xml_path = Path(xml_temp.name) / f"{len(book.worksheets)}.xml"
            descriptor = os.open(xml_path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            os.close(descriptor)
            sheet._writer = _OwnedWorksheetWriter(sheet, out=str(xml_path))
            sheet._writer.write_top()
            sheet.append([encoder.cell(sheet, col[0], header=True) for col in sheet_columns])
        for name, sheet_columns in columns.items():
            sheet = book[name]
            source = explanations if name == "字段说明" else rows.iter_rows(name)
            for row in source:
                _check_deadline(deadline)
                if name == "字段说明" and row["字段或项目"] == "xml_replacement_count":
                    row["说明或值"] = encoder.replacements
                if name == "字段说明":
                    row = dict(row)
                    row["字段或项目"] = METADATA_LABELS.get(row["字段或项目"], row["字段或项目"])
                cells = []
                for title, base, index, count in sheet_columns:
                    value = row.get(base)
                    if isinstance(value, str) and count > 1 and not base.startswith("动作质量详情"):
                        value = value[index * TEXT_CHUNK_SIZE : (index + 1) * TEXT_CHUNK_SIZE]
                    cells.append(encoder.cell(sheet, value, base))
                sheet.append(cells)
        rows.metadata["xml_replacement_count"] = encoder.replacements
        _check_deadline(deadline)
        with ZipFile(
            _DeadlineFile(handle, deadline), "w", ZIP_DEFLATED, allowZip64=True
        ) as archive:
            ExcelWriter(book, archive).save()
        _check_deadline(deadline)
        handle.seek(0)
        return handle
    except BaseException:
        handle.close()
        raise
    finally:
        try:
            if book is not None:
                _cleanup_workbook(book)
        except BaseException:
            handle.close()
            raise
        finally:
            if xml_temp is not None:
                xml_temp.cleanup()
