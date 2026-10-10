from pathlib import Path
import pytest
import json

from django.conf import settings
from docx import Document
from docx.table import _Cell

from apps.crf.completeness import evaluate_completeness
from apps.crf.registry_loader import load_crf_registry
from apps.crf.services.export_docx import export_preview_to_docx
from tests.fixtures.crf_sentinel_cases import clinical_case, slot_marker


def make_preview(baseline=None, visits=None):
    if baseline is None:
        baseline, visits = clinical_case()
    return {
        **evaluate_completeness(load_crf_registry(), baseline, visits or {}),
        "patient": {"name": "测试患者"},
        "project": {"name": "测试项目"},
        "group": {"name": "测试分组"},
    }


def exported(tmp_path, preview=None):
    return Document(export_preview_to_docx(preview or make_preview(), tmp_path / "filled.docx"))


def test_export_fills_original_cells_not_only_appended_summary(tmp_path):
    doc = exported(tmp_path)
    assert "S-DOCX-031" in _Cell(doc.tables[0].rows[7]._tr.tc_lst[1], doc.tables[0]).text
    assert "ZXY" in _Cell(doc.tables[8].rows[2]._tr.tc_lst[1], doc.tables[8]).text


def test_every_source_slot_has_persistent_physical_cell_geometry(tmp_path):
    from apps.crf.services.docx_binding import resolve_source_cells

    manifest = json.loads(
        (Path(settings.CRF_REGISTRY_PATH).parent / "source_slots.v1.json").read_text()
    )
    doc = Document(settings.CRF_TEMPLATE_PATH)
    for position in manifest["slots"]:
        resolve_source_cells(doc, position)[0].paragraphs[0].add_run(
            slot_marker(position["slot_id"])
        )
    path = tmp_path / "geometry.docx"
    doc.save(path)
    saved = Document(path)
    for position in manifest["slots"]:
        assert slot_marker(position["slot_id"]) in resolve_source_cells(saved, position)[0].text
    date_cell = resolve_source_cells(
        saved, {"index_space": "xml_tc", "table_index": 9, "row_index": 2, "cell_index": 1}
    )[0]
    assert "定位_T09_R02_S01" in date_cell.text
    assert date_cell._tc is saved.tables[9].rows[2].cells[2]._tc
    assert date_cell._tc is not saved.tables[9].rows[2].cells[1]._tc


def test_repeat_rows_beyond_original_capacity_are_all_exported(tmp_path):
    from apps.crf.services.docx_binding import resolve_source_cells

    doc = exported(tmp_path)
    assert len(doc.tables[9].rows) == 13
    for table_index, prefix in ((20, "T1药物"), (27, "T2药物"), (21, "T1事件"), (28, "T2事件")):
        text = "\n".join(c.text for row in doc.tables[table_index].rows for c in row.cells)
        for i in range(5):
            assert f"{prefix}{i}" in text
    original_allergy = resolve_source_cells(
        doc, {"index_space": "xml_tc", "table_index": 9, "row_index": 7, "cell_index": 1}
    )[0]
    assert original_allergy.text == "过敏原0"
    assert resolve_source_cells(
        doc, {"index_space": "xml_tc", "table_index": 9, "row_index": 5, "cell_index": 1}
    )[0].text.startswith("□否  ☑是")
    text = "\n".join(c.text for row in doc.tables[9].rows for c in row.cells)
    for i in range(5):
        assert f"手术记录{i}" in text
    for i in range(4):
        assert f"过敏原{i}" in text


def test_hidden_old_entries_are_not_exported_as_active(tmp_path):
    baseline, visits = clinical_case()
    baseline["surgery_allergy"]["has_surgery"] = "否"
    visits["T1"]["form_data"]["crf"]["medication_change"]["has_changes"] = "否"
    doc = exported(tmp_path, make_preview(baseline, visits))
    assert len(doc.tables[9].rows) == 11
    for table_index, hidden in ((9, "手术记录"), (20, "T1药物")):
        assert hidden not in "\n".join(
            c.text for row in doc.tables[table_index].rows for c in row.cells
        )


def physical(doc, table, row, cell):
    from apps.crf.services.docx_binding import resolve_source_cells

    return resolve_source_cells(
        doc, {"index_space": "xml_tc", "table_index": table, "row_index": row, "cell_index": cell}
    )[0]


def test_every_legal_source_binding_is_filled_in_its_own_cell_or_choice_group(tmp_path):
    from apps.crf.registry_loader import load_crf_source_slots
    from apps.crf.services.docx_binding import resolve_source_cells
    from tests.fixtures.crf_sentinel_cases import binding_projection_case

    preview, expected = binding_projection_case(load_crf_registry())
    doc = exported(tmp_path, preview)
    slots = {slot["slot_id"]: slot for slot in load_crf_source_slots()["slots"]}
    seen = set()
    for definition, position, value in expected:
        seen.add(position["slot_id"])
        cells = resolve_source_cells(doc, position)
        slot = slots[position["slot_id"]]
        if slot["options"]:
            mapped = (
                definition.get("export_choice_mapping", {}).get(str(value), {}).get("value", value)
            )
            selected = mapped if isinstance(mapped, list) else [mapped]
            if len(cells) > 1:
                for option, cell in zip(slot["options"], cells):
                    assert cell.text == ("☑" if option in [str(v) for v in selected] else "□"), (
                        position,
                        value,
                        cell.text,
                    )
            else:
                # Distinct same-cell groups get additional literal tests below.
                for option in slot["options"]:
                    symbol = "☑" if option in [str(v) for v in selected] else "□"
                    assert symbol + option in cells[0].text, (position, value, cells[0].text)
        else:
            if definition["widget"] == "date":
                value = (
                    "2024年07月"
                    if definition.get("date_precision") == "month"
                    else "2024年07月09日"
                )
            assert str(value) in cells[0].text, (position, value, cells[0].text)
    from django.conf import settings
    source = Document(settings.CRF_TEMPLATE_PATH)
    static = set()
    for binding in load_crf_registry()['document_bindings']:
        assert binding['kind'] == 'template_static'
        for position in binding['source_positions']:
            static.add(position['slot_id'])
            assert [cell.text for cell in resolve_source_cells(doc,position)] == [cell.text for cell in resolve_source_cells(source,position)]
    assert seen.isdisjoint(static)
    assert seen | static == set(slots)
    for ti, ri in ((0, 7), (8, 1), (3, 0), (15, 0), (22, 0)):
        assert physical(doc, ti, ri, 1).text == "S-SHARED-051"


def test_choices_inside_same_cell_do_not_tick_neighboring_groups(tmp_path):
    baseline, visits = clinical_case()
    visits["T0"]["form_data"]["assessments"]["physical"] = {
        "balance": {
            "side_by_side": {"able": "能", "seconds": 10},
            "semi_tandem": {"able": "不能", "seconds": 3},
            "tandem": {"able": "能", "seconds": 8},
        },
    }
    # Use exact registry paths, with different answers in all three groups.
    fields = load_crf_registry()["fields"]
    selected = [
        f
        for f in fields
        if any(
            p["table_index"] == 13 and p["row_index"] == 2 for p in f.get("source_positions", [])
        )
    ]
    from tests.fixtures.crf_sentinel_cases import binding_projection_case

    preview, _ = binding_projection_case({"fields": selected})
    choices = [
        s["fields"][0] for s in preview["sections"] if s["fields"][0]["widget"] == "single_choice"
    ]
    for field, value in zip(choices, ["能", "不能", "能"]):
        field["value"] = value
    doc = exported(tmp_path, preview)
    assert (
        physical(doc, 13, 2, 1).text
        == "并脚站立10秒：☑能 □不能，实际1秒；半串联站立10秒：□能 ☑不能，实际1秒；串联站立10秒：☑能 □不能，实际1秒；平衡得分1/4"
    )


def test_unknown_na_partial_dates_month_and_ethnicity(tmp_path):
    baseline, visits = clinical_case()
    baseline["demographics"].update(
        {
            "birth_date": {
                "year": {"state": "value", "value": 1950},
                "month": {"state": "unknown"},
                "day": {"state": "value", "value": 9},
            },
            "_crf_field_meta": {
                "patient_baseline.demographics.birth_date": {
                    "state": "value",
                    "reason": "月份遗忘",
                },
                "patient_baseline.demographics.age_years": {
                    "state": "unknown",
                    "reason": "年龄待查",
                },
                "patient_baseline.demographics.gender": {
                    "state": "not_applicable",
                    "reason": "记录不适用",
                },
            },
        }
    )
    baseline["comorbidities"]["cm_coronary"]["diagnosed_at"] = "2020-07"
    doc = exported(tmp_path, make_preview(baseline, visits))
    assert physical(doc, 8, 3, 1).text == "99周岁"
    assert physical(doc, 8, 4, 1).text == "1950年99月09日"
    assert "88" in physical(doc, 8, 5, 1).text
    assert "☑" not in physical(doc, 8, 5, 1).text
    assert physical(doc, 8, 8, 1).text == "□汉族  □黎族  □苗族  ☑其他蒙古族"
    assert physical(doc, 10, 1, 3).text == "2020年07月"
    text = "\n".join(p.text for p in doc.paragraphs)
    for reason in ("月份遗忘", "年龄待查", "记录不适用"):
        assert reason in text


def test_ordinary_missing_unresolved_raw_and_manual_only(tmp_path):
    baseline, visits = clinical_case()
    baseline["demographics"]["age_years"] = "坏年龄旧值"
    visits["T0"]["form_data"]["assessments"] = {}
    visits["T0"]["form_data"]["computed_assessments"] = {"sppb": {"total": 12}}
    baseline["surgery_allergy"]["old_text"] = "旧手术原文待核对"
    doc = exported(tmp_path, make_preview(baseline, visits))
    assert physical(doc, 8, 3, 1).text == "____周岁"
    assert physical(doc, 13, 5, 1).text == "____/12"
    assert physical(doc, 0, 2, 1).text == "________________"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "坏年龄旧值" in text
    assert "旧手术原文待核对" in text


def test_terminal_whole_group_selection_and_conflict_are_traceable(tmp_path):
    baseline, visits = clinical_case()
    visits["T1"]["form_data"]["crf"]["completion_record"] = {
        "status": "提前退出",
        "date": "2026-02-12",
    }
    visits["T2"]["form_data"]["crf"]["completion_record"] = {"status": "完成全部随访"}
    doc = exported(tmp_path, make_preview(baseline, visits))
    assert physical(doc, 29, 0, 1).text.startswith("☑完成全部随访  □提前退出")
    assert physical(doc, 29, 1, 1).text == "|__|__|__|__|年|__|__|月|__|__|日"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "提前退出" in text
    assert "冲突" in text


def test_original_static_cells_paragraphs_and_formatting_survive(tmp_path):
    from apps.crf.registry_loader import load_crf_source_slots
    from docx.oxml.ns import qn

    before = Document(settings.CRF_TEMPLATE_PATH)
    after = exported(tmp_path)
    filled = {
        (s["table_index"], s["row_index"], c)
        for s in load_crf_source_slots()["slots"]
        for c in s["cell_indices"]
    }
    for ti, table in enumerate(before.tables):
        for ri, row in enumerate(table._tbl.tr_lst):
            for ci, tc in enumerate(row.tc_lst):
                original = _Cell(tc, table)
                saved = physical(after, ti, ri, ci)
                assert original._tc.tcPr.xml == saved._tc.tcPr.xml
                if (ti, ri, ci) not in filled:
                    assert original.text == saved.text
    assert [p.text for p in after.paragraphs[: len(before.paragraphs)]] == [
        p.text for p in before.paragraphs
    ]
    names = [b.get(qn("w:name")) for b in after.element.iter(qn("w:bookmarkStart"))]
    ids = [b.get(qn("w:id")) for b in after.element.iter(qn("w:bookmarkStart"))]
    assert len(names) == len(set(names))
    assert len(ids) == len(set(ids))


def test_status_reasons_use_existing_remark_cells_and_preserve_entered_remarks(tmp_path):
    baseline, visits = clinical_case()
    visits["T2"]["form_data"]["crf"]["qc_record"] = {"items": {"item_1": {"remark": "原备注"}}}
    visits["T2"]["form_data"]["crf"]["_field_meta"] = {
        "visit.form_data.crf.qc_record.items.item_1.judgment": {
            "state": "not_applicable",
            "reason": "本例无需复测",
        },
    }
    doc = exported(tmp_path, make_preview(baseline, visits))
    assert physical(doc, 30, 1, 1).text == "□"
    assert physical(doc, 30, 1, 2).text == "□"
    assert physical(doc, 30, 1, 3).text == "☑"
    remark = physical(doc, 30, 1, 4).text
    assert "原备注" in remark and "本例无需复测" in remark and "88" in remark


def test_three_visit_answers_and_all_satisfaction_choices_are_independent(tmp_path):
    doc = exported(tmp_path)
    for ti, expected in ((3, "2026年01月08日"), (15, "2026年02月08日"), (22, "2026年03月08日")):
        assert physical(doc, ti, 1, 1).text == expected
    for ti, expected in ((13, "3/12"), (17, "4/12"), (24, "5/12")):
        assert physical(doc, ti, 5, 1).text == expected
    for table, index, vt in ((19, 1, "T1"), (26, 2, "T2")):
        for row in range(1, 6):
            chosen = (row + index) % 5 + 1
            for column in range(1, 6):
                assert physical(doc, table, row, column).text == ("☑" if column == chosen else "□")
            assert physical(doc, table, row, 6).text == f"{vt}反馈{row}"
        assert physical(doc, table, 6, 6).text == "总分15/25"


def test_bookmarks_respect_paragraph_properties_and_combination_order(tmp_path):
    from docx.oxml.ns import qn

    preview = make_preview()
    for section in preview["sections"]:
        for field in section["fields"]:
            for row in field["rows"]:
                row["fields"].reverse()
    doc = exported(tmp_path, preview)
    assert physical(doc, 21, 2, 6).text == "处理：处理0\n转归：转归0"
    assert list(physical(doc, 21, 2, 6)._tc.iter(qn("w:br")))
    for p in doc.element.iter(qn("w:p")):
        properties = p.find(qn("w:pPr"))
        if properties is not None:
            assert p[0] is properties


def test_missing_reason_keeps_explicit_unknown_status_but_never_exports_residual_raw(tmp_path):
    baseline, visits = clinical_case()
    baseline["demographics"]["_crf_field_meta"] = {
        "patient_baseline.demographics.age_years": {"state": "unknown"}
    }
    doc = exported(tmp_path, make_preview(baseline, visits))
    assert physical(doc, 8, 3, 1).text == "99周岁"
    assert "说明待补" not in "\n".join(p.text for p in doc.paragraphs)


def test_legacy_row_metadata_and_partial_date_conflict_remain_supplemental(tmp_path):
    baseline, visits = clinical_case()
    baseline["surgery_allergy"]["surgery_records"][0]["_field_meta"] = {
        "name": {"state": "value", "old_note": "旧手术注释"},
        "abandoned_field": {"state": "unknown", "reason": "旧原因"},
    }
    visits["T0"]["form_data"]["crf"]["meta"]["visit_date_parts"] = {
        "year": {"state": "value", "value": 2025},
        "month": {"state": "value", "value": 10},
        "day": {"state": "unknown"},
    }
    preview = make_preview(baseline, visits)
    from copy import deepcopy

    before = deepcopy(preview)
    doc = exported(tmp_path, preview)
    assert preview == before
    assert physical(doc, 9, 2, 0).text == "手术记录0"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "旧手术注释" in text and "旧原因" in text
    assert "手术记录 第1行" in text
    assert "2026-01-08" in text and "2025" in text
    assert physical(doc, 3, 1, 1).text == "|__|__|__|__|年|__|__|月|__|__|日"


def test_original_paragraph_properties_stay_first_child(tmp_path):
    from docx.oxml.ns import qn

    for p in exported(tmp_path).element.iter(qn("w:p")):
        properties = p.find(qn("w:pPr"))
        if properties is not None:
            assert p[0] is properties


def test_multiline_combination_uses_word_line_breaks(tmp_path):
    from docx.oxml.ns import qn

    assert list(physical(exported(tmp_path), 21, 2, 6)._tc.iter(qn("w:br")))


def test_unknown_repeat_status_does_not_expand_hidden_old_rows(tmp_path):
    baseline, visits = clinical_case()
    baseline["demographics"]["_crf_field_meta"] = {
        "patient_baseline.surgery_allergy.surgery_records": {"state": "unknown"}
    }
    doc = exported(tmp_path, make_preview(baseline, visits))
    assert len(doc.tables[9].rows) == 11  # only active allergy adds two
    assert "手术记录0" not in "\n".join(c.text for row in doc.tables[9].rows for c in row.cells)


def test_multiline_entered_text_in_existing_placeholder_uses_word_breaks(tmp_path):
    from docx.oxml.ns import qn

    baseline, visits = clinical_case()
    visits["T0"]["form_data"]["crf"]["meta"]["assessor_name"] = "研究者甲\n研究者乙"
    doc = exported(tmp_path, make_preview(baseline, visits))
    cell = physical(doc, 3, 4, 1)
    assert "研究者甲\n研究者乙" in cell.text
    assert list(cell._tc.iter(qn("w:br")))


@pytest.mark.django_db
def test_database_readback_and_preview_export_use_latest_baseline_and_manual_values(
    tmp_path, project_patient, doctor
):
    from apps.patients.models import PatientBaseline
    from apps.visits.models import VisitRecord
    from apps.crf.services.aggregate import build_crf_preview

    baseline, visits = clinical_case()
    PatientBaseline.objects.update_or_create(
        patient=project_patient.patient,
        defaults={**baseline, "created_by": doctor, "updated_by": doctor},
    )
    for vt, payload in visits.items():
        VisitRecord.objects.filter(project_patient=project_patient, visit_type=vt).update(**payload)
    project_patient.refresh_from_db()
    preview = build_crf_preview(project_patient)
    doc = exported(tmp_path, preview)
    assert physical(doc, 0, 7, 1).text == "S-DOCX-031"
    assert physical(doc, 9, 2, 1).text == "2025年06月12日"
    assert physical(doc, 9, 7, 1).text == "过敏原0"
    assert physical(doc, 19, 1, 6).text == "T1反馈1"
    assert physical(doc, 26, 1, 6).text == "T2反馈1"


def test_static_text_inside_filled_cells_and_extension_style_are_preserved(tmp_path):
    import re
    from apps.crf.registry_loader import load_crf_source_slots
    from tests.fixtures.crf_sentinel_cases import binding_projection_case
    from apps.crf.services.docx_binding import resolve_source_cells

    preview, _ = binding_projection_case(load_crf_registry())
    doc = exported(tmp_path, preview)
    for slot in load_crf_source_slots()["slots"]:
        for original, saved in zip(slot["source_cells"], resolve_source_cells(doc, slot)):
            for static in re.split(r"[_□|]+", original):
                if static.strip():
                    assert static in saved.text, (slot["slot_id"], static, saved.text)
    extended = exported(tmp_path)
    original = Document(settings.CRF_TEMPLATE_PATH)
    for table_index, source_row_index, extra_indices in (
        (9, 4, (5, 6)),
        (9, 8, (11, 12)),
        (20, 3, (4, 5, 6)),
        (21, 4, (5, 6)),
    ):
        source_row = original.tables[table_index]._tbl.tr_lst[source_row_index]
        for index in extra_indices:
            row = extended.tables[table_index]._tbl.tr_lst[index]
            assert [tc.tcPr.xml for tc in row.tc_lst] == [tc.tcPr.xml for tc in source_row.tc_lst]
    for table_index in (21, 28):
        assert [
            _Cell(row.tc_lst[0], extended.tables[table_index]).text
            for row in extended.tables[table_index]._tbl.tr_lst[2:]
        ] == ["1", "2", "3", "4", "5"]


@pytest.mark.parametrize("state", ["unknown", "not_applicable"])
def test_inactive_diagnosis_reason_stays_historical_after_parent_changes_to_no(tmp_path, state):
    baseline, visits = clinical_case()
    reason = "既往诊断月份忘记" if state == "unknown" else "既往月份不适用说明"
    baseline["demographics"]["_crf_field_meta"] = {
        "patient_baseline.comorbidities.cm_coronary.diagnosed_at": {
            "state": state,
            "reason": reason,
        }
    }
    active = exported(tmp_path, make_preview(baseline, visits))
    assert reason in physical(active, 10, 1, 4).text
    baseline["comorbidities"]["cm_coronary"]["personal"] = "否"
    preview = make_preview(baseline, visits)
    field = next(
        f
        for s in preview["sections"]
        for f in s["fields"]
        if f["storage"] == "patient_baseline.comorbidities.cm_coronary.diagnosed_at"
    )
    assert not field["applicable"] and field["value"] is None
    saved = exported(tmp_path, preview)
    assert physical(saved, 10, 1, 2).text == "□不详 ☑否 □是"
    assert physical(saved, 10, 1, 3).text == "____年____月"
    assert reason not in physical(saved, 10, 1, 4).text
    historical = [p.text for p in saved.paragraphs if reason in p.text]
    assert len(historical) == 1
    assert "历史非活动" in historical[0]
    assert "2020-07" in historical[0]


@pytest.mark.parametrize("state", ["unknown", "not_applicable"])
def test_ongoing_medication_keeps_old_stop_date_reason_out_of_current_remarks(tmp_path, state):
    baseline, visits = clinical_case()
    rows = visits["T1"]["form_data"]["crf"]["medication_change"]["entries"]
    for index in (0, 4):
        rows[index]["ongoing"] = "仍在使用"
        rows[index]["_field_meta"] = {
            "stopped_on": {"state": state, "reason": f"旧停药日期原因{index}"}
        }
    saved = exported(tmp_path, make_preview(baseline, visits))
    for index, physical_row in ((0, 2), (4, 6)):
        row = saved.tables[20]._tbl.tr_lst[physical_row]
        assert _Cell(row.tc_lst[3], saved.tables[20]).text == "____年__月__日/☑仍在使用"
        assert _Cell(row.tc_lst[5], saved.tables[20]).text == f"药物备注{index}"
        paragraphs = [p.text for p in saved.paragraphs if f"旧停药日期原因{index}" in p.text]
        assert len(paragraphs) == 1
        assert "历史非活动" in paragraphs[0]
        assert f"第{index + 1}行" in paragraphs[0]


def test_supplement_sources_are_clinical_names_but_original_technical_text_is_unchanged(tmp_path):
    baseline, visits = clinical_case()
    baseline["demographics"]["_crf_field_meta"] = {
        "patient_baseline.demographics.age_years": {"state": "unknown", "reason": "年龄待查"}
    }
    baseline["surgery_allergy"]["old_text"] = (
        "旧原文含patient_baseline.demographics与_field_meta字样"
    )
    saved = exported(tmp_path, make_preview(baseline, visits))
    paragraphs = [p.text for p in saved.paragraphs]
    assert "患者基线 人口学信息" in paragraphs
    note = next(p for p in paragraphs if "年龄待查" in p)
    assert "patient_baseline" not in note and "visit.form_data" not in note
    assert not any("来源：patient_baseline" in p or "来源：visit." in p for p in paragraphs)
    assert any("旧原文含patient_baseline.demographics与_field_meta字样" in p for p in paragraphs)


def test_repeat_tables_have_contiguous_header_prefix_without_marking_data_rows(tmp_path):
    from docx.oxml.ns import qn

    source = Document(settings.CRF_TEMPLATE_PATH)
    document = exported(tmp_path)
    for index in (20, 21, 27, 28):
        table = document.tables[index]
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        assert table.rows[1]._tr.trPr.find(qn("w:tblHeader")) is not None
        # Source row identity survives expansion, including T9's second header.
        for row_index, row in enumerate(source.tables[index].rows[2:], 2):
            original_header = row._tr.trPr.find(qn("w:tblHeader")) is not None
            from apps.crf.services.docx_binding import _source_row
            _, saved_row = _source_row(document, {"index_space": "xml_tc", "table_index": index, "row_index": row_index})
            assert (saved_row.trPr.find(qn("w:tblHeader")) is not None) == original_header


def test_multi_section_repeat_table_keeps_original_headers_on_allergy_continuation(tmp_path):
    from docx.oxml.ns import qn
    from apps.crf.services.docx_binding import _source_row

    source = Document(settings.CRF_TEMPLATE_PATH)
    document = exported(tmp_path)
    for row_index, row in enumerate(source.tables[9].rows):
        _, saved = _source_row(document, {"index_space": "xml_tc", "table_index": 9, "row_index": row_index})
        assert (saved.trPr.find(qn("w:tblHeader")) is not None) == (row._tr.trPr.find(qn("w:tblHeader")) is not None)


@pytest.mark.parametrize("state", ["unknown", "not_applicable"])
@pytest.mark.parametrize("value", [
    {"year": None, "month": {"state": "value", "value": 13}, "day": None},
    {"year": 1960, "month": {"state": "unknown"}, "day": None},
    {"year": {"state": "unknown", "value": 1960}, "month": None, "day": None},
])
def test_invalid_historical_date_metadata_never_reauthorizes_original_cell(tmp_path, state, value):
    baseline = {"demographics": {"birth_date": value, "_crf_field_meta": {
        "patient_baseline.demographics.birth_date": {"state": state, "reason": "历史日期待核对"}
    }}}
    preview = make_preview(baseline, {})
    item = next(f for s in preview["sections"] for f in s["fields"] if f["field_id"] == "dm_birth_date")
    assert item["state"] == "unresolved" and item["issues"]
    doc = exported(tmp_path, preview)
    from apps.crf.services.docx_binding import resolve_source_cells
    definition = next(f for f in load_crf_registry()["fields"] if f["field_id"] == "dm_birth_date")
    cell = resolve_source_cells(doc, definition["source_positions"][0])[0].text
    assert "13月" not in cell and "1960年" not in cell and "9999年" not in cell
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "历史日期待核对" in text
    assert json.dumps(value, ensure_ascii=False) in text


def test_valid_observation_with_extra_metadata_retains_history_in_preview_and_export(tmp_path):
    baseline = {"surgery_allergy": {"has_surgery": "是", "surgery_records": [{"_row_id": "a", "name": "旧手术"}]}, "demographics": {"_crf_field_meta": {
        "patient_baseline.surgery_allergy.has_surgery": {"state": "value", "legacy_feedback": "旧反馈必须保留"}
    }}}
    preview = make_preview(baseline, {})
    item = next(f for s in preview["sections"] for f in s["fields"] if f["field_id"] == "sa_has_surgery")
    assert item["state"] == "normal" and item["metadata"]["legacy_feedback"] == "旧反馈必须保留"
    doc = exported(tmp_path, preview)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "旧手术" in text
    assert "旧反馈必须保留" in text
