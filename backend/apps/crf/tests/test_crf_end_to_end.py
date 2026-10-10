"""Release blockers and genuine API → DOCX → LibreOffice PDF acceptance."""

import json
from pathlib import Path
import shutil
import subprocess
import sys

import pytest

ROOT = Path(__file__).resolve().parents[4]
REGISTRY = Path("specs/patient-rehab-system/crf/registry.v1.json")
TEMPLATE = Path("docs/other/认知衰弱数字疗法研究_CRF表_修订稿.docx")


def release(root):
    return subprocess.run(
        [sys.executable, str(ROOT / "scripts/verify_crf_release.py"), "--repo-root", str(root)],
        capture_output=True,
        text=True,
        check=False,
    )


def test_release_accepts_locked_environment_and_all_source_positions():
    process = release(ROOT)
    assert process.returncode == 0, process.stdout + process.stderr
    result = json.loads(process.stdout)
    assert result["source_coverage"] == {"tables": 31, "rows": 265, "slots": 561}
    assert result["contract_errors"] == []
    assert result["schema_revision"] == 3
    assert result["template_revision"] == "1.1"


@pytest.mark.parametrize(
    "damage",
    [
        "missing_asset",
        "different_copy",
        "changed_template",
        "unlocked",
        "wrong_installed_version",
        "missing_binding",
        "missing_extra",
    ],
)
def test_release_cli_fails_closed_for_damaged_release(tmp_path, damage):
    for relative in [
        REGISTRY,
        TEMPLATE,
        REGISTRY.with_name("source_slots.v1.json"),
        Path("frontend/src/crf/registry.v1.json"),
        Path("backend/requirements.lock.txt"),
        Path("backend/requirements-dev.lock.txt"),
    ]:
        target = tmp_path / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT / relative, target)
    if damage == "missing_asset":
        (tmp_path / REGISTRY).unlink()
    elif damage == "different_copy":
        (tmp_path / "frontend/src/crf/registry.v1.json").write_text("{}")
    elif damage == "changed_template":
        with (tmp_path / TEMPLATE).open("ab") as stream:
            stream.write(b"tamper")
    elif damage == "missing_extra":
        lock = tmp_path / "backend/requirements.lock.txt"
        lock.write_text(
            "\n".join(
                line
                for line in lock.read_text().splitlines()
                if not line.startswith("psycopg-binary==")
            )
            + "\n"
        )
    elif damage in {"unlocked", "wrong_installed_version"}:
        lock = tmp_path / "backend/requirements.lock.txt"
        lock.write_text("Django>=5.0\n" if damage == "unlocked" else "Django==0.0.0\n")
    else:
        registry = json.loads((tmp_path / REGISTRY).read_text())
        victim = next(f for f in registry["fields"] if f.get("source_positions"))
        victim["source_positions"] = []
        for relative in [REGISTRY, Path("frontend/src/crf/registry.v1.json")]:
            (tmp_path / relative).write_text(json.dumps(registry))
    process = release(tmp_path)
    assert process.returncode != 0
    assert "CRF 发布校验失败" in process.stderr


def _value(field):
    if field["widget"] == "number":
        return max(0, field.get("min", field.get("validation", {}).get("min", 0)))
    if field["widget"] == "date":
        return "2026-10" if field.get("date_precision") == "month" else "2026-10-08"
    if field["widget"] == "single_choice":
        return field["options"][0]
    if field["widget"] == "multi_choice":
        return [field["options"][0]]
    if field["widget"] == "table_repeat":
        return [
            {"_row_id": f"row-{i}", **{c["storage"]: _value(c) for c in field["children"]}}
            for i in range(5)
        ]
    return "已核对"


def _activate(condition, context):
    from tests.fixtures.crf_sentinel_cases import put

    for child in condition.get("all", []):
        _activate(child, context)
    if "storage" not in condition:
        return
    if condition["op"] == "eq":
        put(context, condition["storage"], condition["value"])
    elif condition["op"] == "in":
        put(context, condition["storage"], condition["value"][0])
    elif condition["op"] == "contains":
        put(context, condition["storage"], [condition["value"]])


def api_case(kind):
    from apps.crf.registry_loader import load_crf_registry
    from tests.fixtures.crf_sentinel_cases import complete_known_case, put

    baseline, visits = complete_known_case()
    if kind == "partial":
        return {"subject_id": "部分病例", "demographics": {"age_years": 0}}, {
            "T0": {"visit_date": None, "form_data": {"assessments": {"sppb": {"balance": 0}}}}
        }
    for field in load_crf_registry()["fields"]:
        if field.get("legacy") or field.get("canonical_storage"):
            continue
        for vt in field.get("visit_types") or [None]:
            context = {"patient_baseline": baseline, "visit": visits.get(vt)}
            if not field["include_in_completeness"] or field["widget"] == "table_repeat":
                put(context, field["storage"], _value(field))
    # Activate compatible dependent groups; mutually exclusive groups are covered
    # by separate clinical cases, never by a projection submitted as an API payload.
    for field in load_crf_registry()["fields"]:
        if field.get("legacy") or field.get("canonical_storage"):
            continue
        for vt in field.get("visit_types") or [None]:
            _activate(
                field.get("visible_when") or {},
                {"patient_baseline": baseline, "visit": visits.get(vt)},
            )
    baseline["subject_id"] = "E2E-" + kind
    baseline["name_initials"] = "ZXY"
    baseline["demographics"]["ethnicity"] = "其他"
    baseline["demographics"]["ethnicity_other_remark"] = "蒙古族"
    baseline["lifestyle"]["smoking_status"] = "目前吸烟" if kind == "current" else "已戒烟"
    baseline["lifestyle"]["drinking_status"] = "目前饮酒" if kind == "current" else "已戒酒"
    for vt, visit in visits.items():
        a = visit["form_data"]["assessments"]
        a["sppb"].update(balance=4, gait=3, chair_stand=2, total=9)
        a["moca"]["total"] = 1  # 0 in each of seven domains + education bonus 1.
        a["moca"]["leq22"] = "是"
        if kind == "current":
            a["physical"]["chair"]["unable"] = "能完成5次"
            a["physical"]["tug"]["unable"] = "能完成"
        crf = visit["form_data"]["crf"]
        if vt != "T0":
            for item in crf["satisfaction"]["items"].values():
                item["score"] = 1
                item["remark"] = f"{vt}独立反馈"
            crf["satisfaction"]["total"] = 5
        if vt != "T2":
            crf.pop("completion_record", None)
            crf.pop("qc_record", None)
    completion = visits["T2"]["form_data"]["crf"]["completion_record"]
    completion["status"] = {"current": "完成全部随访", "former": "其他"}.get(kind, "提前退出")
    if kind == "unknown_na":
        baseline["demographics"]["_crf_field_meta"] = {
            "patient_baseline.demographics.age_years": {
                "state": "unknown",
                "reason": "原始年龄未提供",
            },
            "patient_baseline.demographics.education_years": {
                "state": "not_applicable",
                "reason": "无法按学年换算",
            },
        }
        # Partial date has independent known/unknown components and field reason.
        baseline["comorbidities"]["cm_coronary"]["diagnosed_at"] = {
            "year": {"state": "value", "value": 2020},
            "month": {"state": "unknown"},
        }
        baseline["demographics"]["_crf_field_meta"].update(
            {
                "patient_baseline.comorbidities.cm_coronary.diagnosed_at": {
                    "state": "value",
                    "reason": "仅有年份",
                }
            }
        )
    if kind == "unknown_na":
        row = baseline["surgery_allergy"]["surgery_records"][0]
        row["_field_meta"] = {"hospital_note": {"state": "unknown", "reason": "医院待核实"}}
    return baseline, visits


def _assert_subset(expected, actual):
    if isinstance(expected, dict):
        for key, value in expected.items():
            assert key in actual, (key, actual)
            _assert_subset(value, actual[key])
    else:
        assert actual == expected


def _save_case(client, project_patient, kind):
    from apps.visits.models import VisitRecord

    baseline, visits = api_case(kind)
    url = f"/api/patients/{project_patient.patient_id}/baseline/"
    response = client.patch(url, baseline, format="json")
    assert response.status_code == 200, response.data
    _assert_subset(baseline, client.get(url).data)
    for vt, payload in visits.items():
        visit = VisitRecord.objects.get(project_patient=project_patient, visit_type=vt)
        url = f"/api/visits/{visit.id}/"
        response = client.patch(url, payload, format="json")
        assert response.status_code == 200, (vt, response.data)
        _assert_subset(payload, client.get(url).data)
    if kind == "unknown_na":
        rows = baseline["surgery_allergy"]["surgery_records"]
        reordered = list(reversed(rows))
        url = f"/api/patients/{project_patient.patient_id}/baseline/"
        response = client.patch(
            url, {"surgery_allergy": {"surgery_records": reordered}}, format="json"
        )
        assert response.status_code == 200, response.data
        saved = client.get(url).data["surgery_allergy"]["surgery_records"]
        assert saved == reordered
        assert saved[-1]["_row_id"] == "row-0"
        assert saved[-1]["_field_meta"]["hospital_note"]["reason"] == "医院待核实"
    response = client.get(f"/api/crf/project-patients/{project_patient.id}/preview/")
    assert response.status_code == 200, response.data
    return response.data


@pytest.mark.django_db
@pytest.mark.parametrize("kind", ["current", "former", "early_exit", "partial", "unknown_na"])
def test_real_api_roundtrip_original_docx_and_pdf(
    doctor, project_patient, settings, tmp_path, kind
):
    import os
    from docx import Document
    from pypdf import PdfReader
    from rest_framework.test import APIClient
    from apps.crf.models import CrfExport
    from apps.crf.registry_loader import load_crf_source_slots
    from apps.crf.services.docx_binding import resolve_source_cells

    assert shutil.which(settings.CRF_LIBREOFFICE_BINARY), "真实 LibreOffice 是验收必需依赖"
    directory = Path(os.environ.get("CRF_E2E_ARTIFACT_DIR", tmp_path)) / kind
    directory.mkdir(parents=True, exist_ok=True)
    settings.MEDIA_ROOT = directory
    settings.CRF_EXPORT_DIR = directory / "exports"
    client = APIClient()
    client.force_authenticate(doctor)
    preview = _save_case(client, project_patient, kind)
    if kind in {"current", "former", "early_exit"}:
        assert preview["counts"]["unresolved"] == 0, preview["missing_fields"]
        assert preview["counts"]["unknown"] == preview["counts"]["not_applicable"] == 0
        assert not preview["issues"], preview["issues"]
        for visit in preview["visits"].values():
            scores = visit["form_data"]["assessments"]
            assert scores["sppb"]["total"] == 4 + 3 + 2 == 9
            assert scores["moca"]["total"] == 0 + 0 + 0 + 0 + 0 + 0 + 0 + 1 == 1
    if kind == "unknown_na":
        assert preview["counts"]["unknown"] == 3
        assert preview["counts"]["not_applicable"] == 1
        assert preview["counts"]["unresolved"] == 0
        assert preview["issues"] == []
    if kind == "partial":
        assert preview["counts"]["unresolved"] > 0
        assert preview["visits"]["T0"]["form_data"]["assessments"]["sppb"]["balance"] == 0
    response = client.post(f"/api/crf/project-patients/{project_patient.id}/export/")
    assert response.status_code == 200, response.data
    record = CrfExport.objects.latest("id")
    docx_path, pdf_path = Path(record.docx_file.path), Path(record.pdf_file.path)
    document = Document(docx_path)
    # All 561 physical locations survive reloading the actually exported document.
    positions = load_crf_source_slots()["slots"]
    geometry = {
        p["slot_id"]: [c.text for c in resolve_source_cells(document, p)] for p in positions
    }
    assert len(geometry) == 561
    assert ("部分病例" if kind == "partial" else "E2E-" + kind) in geometry["T00.R07.S00"][0]
    coverage = (
        _check_source_answers(document, preview)
        if kind in {"current", "former", "early_exit"}
        else {}
    )
    (directory / "coverage.json").write_text(json.dumps(coverage, ensure_ascii=False, indent=2))
    pages = [p.extract_text() for p in PdfReader(pdf_path).pages]
    assert len(pages) > 1
    import re
    for index, page_text in enumerate(pages, 1):
        normalized = re.sub(r"\s+", "", page_text)
        assert f"第{index}页/共{len(pages)}页" in normalized
    assert all(text.strip() for text in pages)
    assert all(any("\u4e00" <= char <= "\u9fff" for char in text) for text in pages)
    text = "\n".join(pages)
    assert "衰弱" in text and "Case Report Form" in text
    if kind != "partial":
        assert "蒙古族" in text and "独立反馈" in text
    (directory / "preview.json").write_text(json.dumps(preview, ensure_ascii=False, indent=2))
    (directory / "source-cell-text.json").write_text(
        json.dumps(geometry, ensure_ascii=False, indent=2)
    )
    (directory / "pdf-pages.json").write_text(json.dumps(pages, ensure_ascii=False, indent=2))
    (directory / "export.json").write_text(json.dumps(response.data, ensure_ascii=False, indent=2))


def _leaves(preview):
    for section in preview["sections"]:
        for field in section["fields"]:
            if field["widget"] == "table_repeat":
                for row in field["rows"]:
                    for child in row["fields"]:
                        if field["applicable"]:
                            yield child, row["index"]
            else:
                yield field, None


def _check_source_answers(document, preview):
    from apps.crf.registry_loader import load_crf_registry, load_crf_source_slots
    from apps.crf.services.docx_binding import resolve_source_cells

    definitions = {}
    for field in load_crf_registry()["fields"]:
        definitions[field["field_id"]] = field
        definitions.update({child["field_id"]: child for child in field.get("children", [])})
    manifest = {p["slot_id"]: p for p in load_crf_source_slots()["slots"]}
    seen, fields = set(), set()
    for field, index in _leaves(preview):
        if not field["applicable"] or field["state"] != "normal":
            continue
        definition = definitions[field["field_id"]]
        fields.add(field["field_id"])
        for position in field["source_positions"]:
            if index is not None and position.get("repeat_index") != index:
                continue
            cells = resolve_source_cells(document, position)
            options = manifest[position["slot_id"]]["options"]
            value = field["value"]
            if options:
                mapped = (
                    definition.get("export_choice_mapping", {})
                    .get(str(value), {})
                    .get("value", value)
                )
                selected = [str(v) for v in mapped] if isinstance(mapped, list) else [str(mapped)]
                if len(cells) > 1:
                    for option, cell in zip(options, cells):
                        assert cell.text == ("☑" if option in selected else "□"), position
                else:
                    for option in options:
                        assert ("☑" if option in selected else "□") + option in cells[0].text, (
                            field["field_id"],
                            position,
                            value,
                            cells[0].text,
                        )
            else:
                if field["widget"] == "date":
                    parts = str(value).split("-")
                    value = f"{parts[0]}年{parts[1]}月" + (
                        f"{parts[2]}日" if len(parts) == 3 else ""
                    )
                assert str(value) in cells[0].text, (
                    field["field_id"],
                    position,
                    value,
                    cells[0].text,
                )
            seen.add(position["slot_id"])
    return {"source_slots": sorted(seen), "field_ids": sorted(fields)}


@pytest.mark.django_db
def test_api_case_set_covers_every_current_field_and_all_561_source_positions(
    doctor, project_patient, tmp_path
):
    from docx import Document
    from rest_framework.test import APIClient
    from apps.crf.registry_loader import load_crf_registry, load_crf_source_slots
    from apps.crf.services.export_docx import export_preview_to_docx

    client = APIClient()
    client.force_authenticate(doctor)
    slots, fields = set(), set()
    for kind in ("current", "former", "early_exit"):
        preview = _save_case(client, project_patient, kind)
        document = Document(export_preview_to_docx(preview, tmp_path / f"{kind}.docx"))
        coverage = _check_source_answers(document, preview)
        slots.update(coverage["source_slots"])
        fields.update(coverage["field_ids"])
    expected_fields = set()
    for field in load_crf_registry()["fields"]:
        if field.get("legacy") or field.get("canonical_storage"):
            continue
        expected_fields.update(c["field_id"] for c in field.get("children", [field]))
    assert fields >= expected_fields, sorted(expected_fields - fields)
    static={p['slot_id'] for b in load_crf_registry()['document_bindings'] for p in b['source_positions']}
    assert static=={'T01.R01.S00','T01.R02.S00'}
    assert slots|static == {p["slot_id"] for p in load_crf_source_slots()["slots"]}


@pytest.mark.django_db
def test_legacy_values_unrelated_patch_null_zero_and_row_metadata_remain_exact(
    doctor, project_patient, settings, tmp_path
):
    from copy import deepcopy
    from rest_framework.test import APIClient
    from apps.patients.models import PatientBaseline
    from apps.visits.models import VisitRecord
    from apps.crf.models import CrfExport
    from pypdf import PdfReader

    client = APIClient()
    client.force_authenticate(doctor)
    # Direct model write is solely the pre-existing historical state. Every new
    # user change and the export below uses real API validation.
    old = {"age_years": "旧年龄待核实", "education_years": 9}
    PatientBaseline.objects.create(patient=project_patient.patient, demographics=deepcopy(old))
    visit = VisitRecord.objects.get(project_patient=project_patient, visit_type="T0")
    visit.form_data = {
        "assessments": {"sppb": {"balance": "旧分值"}},
        "crf": {"completion": "历史结束记录 LEGACY-COMPLETION"},
    }
    visit.save()
    url = f"/api/patients/{project_patient.patient_id}/baseline/"
    assert client.patch(url, {"name_initials": "LEG"}, format="json").status_code == 200
    assert client.get(url).data["demographics"] == old
    url = f"/api/visits/{visit.id}/"
    patch = {
        "form_data": {
            "crf": {"meta": {"assessor_name": "历史核对"}},
            "assessments": {"grip_strength_kg": 0},
        }
    }
    response = client.patch(url, patch, format="json")
    assert response.status_code == 200, response.data
    read = client.get(url).data["form_data"]
    assert read["assessments"]["sppb"]["balance"] == "旧分值"
    assert read["crf"]["completion"] == "历史结束记录 LEGACY-COMPLETION"
    assert read["assessments"]["grip_strength_kg"] == 0
    assert (
        client.patch(
            url, {"form_data": {"assessments": {"grip_strength_kg": None}}}, format="json"
        ).status_code
        == 200
    )
    assert client.get(url).data["form_data"]["assessments"]["grip_strength_kg"] is None
    rejected = client.patch(
        url, {"form_data": {"crf": {"completion": {"status": "other"}}}}, format="json"
    )
    assert rejected.status_code == 400
    assert (
        client.get(url).data["form_data"]["crf"]["completion"] == "历史结束记录 LEGACY-COMPLETION"
    )
    import os

    directory = Path(os.environ.get("CRF_E2E_ARTIFACT_DIR", tmp_path)) / "legacy"
    directory.mkdir(parents=True, exist_ok=True)
    settings.MEDIA_ROOT = directory
    settings.CRF_EXPORT_DIR = directory / "exports"
    exported = client.post(f"/api/crf/project-patients/{project_patient.id}/export/")
    assert exported.status_code == 200, exported.data
    text = "\n".join(
        p.extract_text() for p in PdfReader(CrfExport.objects.latest("id").pdf_file.path).pages
    )
    assert "旧分值" in text and "LEGACY-COMPLETION" in text
    from docx import Document

    record = CrfExport.objects.latest("id")
    document = Document(record.docx_file.path)
    assert "历史结束记录 LEGACY-COMPLETION" in "\n".join(p.text for p in document.paragraphs)
    preview = client.get(f"/api/crf/project-patients/{project_patient.id}/preview/").data
    assert preview["issues"]
    (directory / "preview.json").write_text(json.dumps(preview, ensure_ascii=False, indent=2))
    (directory / "export.json").write_text(json.dumps(exported.data, ensure_ascii=False, indent=2))


@pytest.mark.skipif(sys.platform != "linux", reason="Regression targets the deployed Linux LibreOffice pagination")
@pytest.mark.parametrize("with_events", [False, True])
def test_linux_long_feedback_keeps_t2_event_table_after_page_break(tmp_path, with_events):
    from pypdf import PdfReader
    from apps.crf.completeness import evaluate_completeness
    from apps.crf.registry_loader import load_crf_registry
    from apps.crf.services.export_docx import export_preview_to_docx
    from apps.crf.services.export_pdf import convert_docx_to_pdf

    labels = ["操作简单、易学习", "内容清楚、容易理解", "提醒/反馈有助坚持训练", "不适或负担较小", "愿意继续使用或推荐"]
    names = ["operation", "clarity", "reminders", "burden", "recommendation"]
    crf = {
        "meta": {"actual_week": 48, "assessor_name": "UI-T2评估者"},
        "satisfaction": {"items": {name: {"score": 5-i, "remark": f"UI-T2独立反馈{i+1}：{label}。 第二行：研究员现场核对，尾标记T2-{i+1}"} for i, (name, label) in enumerate(zip(names, labels))}, "total": 15, "overall_score": 5},
        "completion_record": {"status": "完成全部随访", "include_analysis": "待统计方案判定"},
        "qc_record": {"items": {"item_1": {"judgment": "否", "remark": "UI质控：本地验收保留其他字段缺失，确认导出可见。"}}, "investigator_name": "UI研究者", "reviewer_name": "UI质控员"},
    }
    if with_events:
        _, visits = api_case("current")
        crf["adverse_events"] = visits["T2"]["form_data"]["crf"]["adverse_events"]
        for i, row in enumerate(crf["adverse_events"]["entries"]):
            row["event_name"] = f"分页事件{i+1}"
    baseline = {}
    if with_events:
        baseline = {"surgery_allergy": {
            "has_surgery": "是", "has_allergy": "是",
            "surgery_records": [{"name": "长备注手术", "hospital_note": "研究员复核手术经过及术后恢复情况，保持多行长备注在原格。" * 10 + "手术长文尾"}],
            "allergy_records": [{"category": "药物", "allergen": f"明确过敏原{i}", "manifestation": "皮疹", "severity_treatment": "已记录"} for i in range(1, 6)],
        }}
    projection = evaluate_completeness(load_crf_registry(), baseline, {"T2": {"visit_date": None, "form_data": {"crf": crf}}})
    source = export_preview_to_docx(projection, tmp_path / "long-feedback.docx")
    pdf = convert_docx_to_pdf(source, tmp_path / "pdf")
    text = "".join("".join(page.extract_text().split()) for page in PdfReader(pdf).pages)
    assert text.count("事件名称/表现") >= 2, "Both T1 and T2 event-table headers must remain visible"
    for i in range(1, 6):
        assert f"尾标记T2-{i}" in text
        if with_events:
            assert f"分页事件{i}" in text
            assert f"明确过敏原{i}" in text
    if with_events:
        assert "手术长文尾" in text


@pytest.mark.django_db
@pytest.mark.parametrize("state", ["unknown", "not_applicable"])
def test_historical_invalid_date_and_extra_metadata_real_api_dual_export(doctor, project_patient, settings, tmp_path, state):
    import os
    from docx import Document
    from pypdf import PdfReader
    from rest_framework.test import APIClient
    from apps.patients.models import PatientBaseline
    from apps.crf.models import CrfExport

    directory = Path(os.environ.get("CRF_E2E_ARTIFACT_DIR", tmp_path)) / ("history_" + state)
    directory.mkdir(parents=True, exist_ok=True)
    settings.MEDIA_ROOT = directory
    settings.CRF_EXPORT_DIR = directory / "exports"
    old = {"year": None, "month": {"state": "value", "value": 13}, "day": None}
    baseline, _ = PatientBaseline.objects.get_or_create(patient=project_patient.patient)
    baseline.demographics = {"birth_date": old, "age_years": 81, "_crf_field_meta": {
        "patient_baseline.demographics.birth_date": {"state": state, "reason": "历史日期待核对"},
        "patient_baseline.surgery_allergy.has_surgery": {"state": "value", "legacy_feedback": "旧反馈必须保留"},
    }}
    baseline.surgery_allergy = {"has_surgery": "是", "surgery_records": [
        {"_row_id": "old-a", "name": "历史手术甲", "hospital_note": "甲医院反馈"},
        {"_row_id": "old-b", "name": "历史手术乙", "_field_meta": {"date": {"state": "unknown", "reason": "乙日期不详"}}},
    ]}
    baseline.save()
    client = APIClient()
    client.force_authenticate(doctor)
    url = f"/api/patients/{project_patient.patient_id}/baseline/"
    response = client.patch(url, {"name_initials": "HISTORY"}, format="json")
    assert response.status_code == 200, response.data
    assert client.get(url).data["demographics"] == baseline.demographics
    response = client.patch(url, {"demographics": {"birth_date": {"month": {"state": "value", "value": 14}}}}, format="json")
    assert response.status_code == 400
    assert client.get(url).data["demographics"] == baseline.demographics
    preview = client.get(f"/api/crf/project-patients/{project_patient.id}/preview/").data
    birth = next(f for s in preview["sections"] for f in s["fields"] if f["field_id"] == "dm_birth_date")
    assert birth["state"] == "unresolved" and birth["raw_value"] == old
    response = client.post(f"/api/crf/project-patients/{project_patient.id}/export/")
    assert response.status_code == 200, response.data
    record = CrfExport.objects.latest("id")
    docx = Path(record.docx_file.path)
    pdf = Path(record.pdf_file.path)
    shutil.copyfile(docx, directory / "filled.docx")
    shutil.copyfile(pdf, directory / "filled.pdf")
    doc = Document(docx)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert json.dumps(birth["raw_value"], ensure_ascii=False) in text
    for marker in ["历史日期待核对", "旧反馈必须保留", "历史手术甲", "甲医院反馈", "乙日期不详"]:
        assert marker in text
        assert marker in "".join(p.extract_text() for p in PdfReader(pdf).pages).replace("\n", "")
    (directory / "preview.json").write_text(json.dumps(preview, ensure_ascii=False, indent=2))


@pytest.mark.django_db
def test_single_answer_acceptance_round_trip(doctor,project_patient,settings,tmp_path):
    from rest_framework.test import APIClient
    from apps.visits.models import VisitRecord
    from apps.crf.services.aggregate import build_crf_preview
    from apps.crf.services.export_docx import export_preview_to_docx
    from apps.crf.services.export_pdf import convert_docx_to_pdf
    from docx import Document
    from pypdf import PdfReader
    client=APIClient();client.force_authenticate(doctor)
    _save_case(client,project_patient,'current')
    baseline_url=f'/api/patients/{project_patient.patient_id}/baseline/'
    marked={'__crf_answer__':1,'state':'unknown','value':0,'reason':None}
    response=client.patch(baseline_url,{'demographics':{'age_years':marked}},format='json')
    assert response.status_code==200,response.data
    assert client.get(baseline_url).data['demographics']['age_years']==marked
    visit=VisitRecord.objects.get(project_patient=project_patient,visit_type='T1');url=f'/api/visits/{visit.pk}/'
    date={'__crf_answer__':1,'state':'not_applicable','value':'2026-10-10','reason':None}
    feedback={'items':{'operation':{'remark':'单字段验收反馈'}}}
    response=client.patch(url,{'visit_date':date,'form_data':{'crf':{'satisfaction':feedback}}},format='json')
    assert response.status_code==200,response.data
    visit.refresh_from_db();assert visit.visit_date is None
    assert visit.form_data['crf']['meta']['visit_date_answer']==date
    assert client.get(url).data['visit_date']==date
    preview=build_crf_preview(project_patient)
    fields=[f for s in preview['sections'] for f in s['fields']]
    age=next(f for f in fields if f['field_id']=='dm_age_years')
    logical=next(f for f in fields if f['field_id']=='vd_T1_visit_date')
    assert age['state']=='unknown' and age['value'] is None and age['raw_value']==marked
    assert logical['state']=='not_applicable' and logical['raw_value']==date
    docx=export_preview_to_docx(preview,tmp_path/'single-answer.docx')
    document=Document(docx)
    # Document revision page remains the unchanged template content.
    source=Document(settings.CRF_TEMPLATE_PATH)
    assert [c.text for row in document.tables[1].rows for c in row.cells]==[c.text for row in source.tables[1].rows for c in row.cells]
    assert '99' in document.tables[8].rows[3].cells[1].text
    assert '88' in document.tables[15].rows[1].cells[1].text
    pdf=convert_docx_to_pdf(docx,tmp_path/'pdf')
    text='\n'.join(page.extract_text() or '' for page in PdfReader(pdf).pages)
    assert '单字段验收反馈' in text and '99' in text and '88' in text
    for answer in ['2026','2026-10','2026-10-10',None]:
        response=client.patch(url,{'visit_date':answer},format='json');assert response.status_code==200,response.data
        visit.refresh_from_db();assert (visit.visit_date.isoformat() if visit.visit_date else None)==('2026-10-10' if answer=='2026-10-10' else None)
        assert client.get(url).data['visit_date']==answer
    response=client.patch(baseline_url,{'demographics':{'age_years':0}},format='json');assert response.status_code==200
    assert client.get(baseline_url).data['demographics']['age_years']==0
