from copy import deepcopy
import pytest

from apps.crf.services.aggregate import build_crf_preview
from apps.patients.models import PatientBaseline
from apps.visits.models import VisitRecord


@pytest.mark.django_db
def test_crf_preview_reports_missing_visit_fields(project_patient):
    preview = build_crf_preview(project_patient)

    assert preview["project_patient_id"] == project_patient.id
    assert "T0.访视日期" in preview["missing_fields"]
    assert "T1.访视日期" in preview["missing_fields"]
    assert "T2.访视日期" in preview["missing_fields"]


@pytest.mark.django_db
def test_crf_preview_snapshot_has_no_wearable_or_health_sections(project_patient):
    preview = build_crf_preview(project_patient)

    assert set(preview) == {
        "project_patient_id",
        "patient",
        "patient_baseline",
        "project",
        "group",
        "visits",
        "missing_fields",
        "sections",
        "issues",
        "counts",
        "schema_revision",
        "template_revision",
    }


@pytest.mark.django_db
def test_crf_preview_reports_missing_visit_assessment_fields(project_patient):
    preview = build_crf_preview(project_patient)
    for vt in ("T0", "T1", "T2"):
        assert f"{vt}.SPPB总分" in preview["missing_fields"]
        assert f"{vt}.MoCA总分" in preview["missing_fields"]


@pytest.mark.django_db
def test_crf_preview_clears_assessment_missing_after_filled(project_patient):
    for vt in ("T0", "T1", "T2"):
        v = VisitRecord.objects.get(project_patient=project_patient, visit_type=vt)
        v.form_data = {
            "assessments": {"sppb": {"total": 9}, "moca": {"total": 22}},
            "computed_assessments": {},
        }
        v.save(update_fields=["form_data"])

    preview = build_crf_preview(project_patient)
    for vt in ("T0", "T1", "T2"):
        assert f"{vt}.SPPB总分" not in preview["missing_fields"]
        assert f"{vt}.MoCA总分" not in preview["missing_fields"]


@pytest.mark.django_db
def test_crf_preview_treats_zero_as_present(project_patient):
    v = VisitRecord.objects.get(project_patient=project_patient, visit_type="T0")
    v.form_data = {
        "assessments": {"sppb": {"total": 0}, "moca": {"total": 0}},
        "computed_assessments": {},
    }
    v.save(update_fields=["form_data"])

    preview = build_crf_preview(project_patient)
    assert "T0.SPPB总分" not in preview["missing_fields"]
    assert "T0.MoCA总分" not in preview["missing_fields"]


@pytest.mark.django_db
def test_crf_preview_returns_empty_patient_baseline_when_absent(project_patient):
    preview = build_crf_preview(project_patient)

    assert "patient_baseline" in preview
    assert preview["patient_baseline"]["subject_id"] == ""
    assert preview["patient_baseline"]["demographics"] == {}

    assert "patient_baseline.受试者编号" in preview["missing_fields"]
    assert "patient_baseline.教育年限（年）" in preview["missing_fields"]


@pytest.mark.django_db
def test_crf_preview_reports_missing_patient_baseline_fields(project_patient, doctor):
    baseline, _ = PatientBaseline.objects.get_or_create(
        patient=project_patient.patient,
        defaults={"created_by": doctor, "updated_by": doctor},
    )

    preview = build_crf_preview(project_patient)
    assert "patient_baseline.受试者编号" in preview["missing_fields"]
    assert "patient_baseline.教育年限（年）" in preview["missing_fields"]

    baseline.subject_id = "S-0001"
    baseline.demographics = {"education_years": 9}
    baseline.updated_by = doctor
    baseline.save()

    preview2 = build_crf_preview(project_patient)
    assert "patient_baseline.受试者编号" not in preview2["missing_fields"]
    assert "patient_baseline.教育年限（年）" not in preview2["missing_fields"]


@pytest.mark.django_db
def test_crf_preview_returns_empty_patient_baseline_when_present_but_all_none(
    project_patient, doctor
):
    baseline, _ = PatientBaseline.objects.get_or_create(
        patient=project_patient.patient,
        defaults={"created_by": doctor, "updated_by": doctor},
    )

    baseline.subject_id = None
    baseline.name_initials = None
    baseline.demographics = None
    baseline.surgery_allergy = None
    baseline.comorbidities = None
    baseline.lifestyle = None
    baseline.baseline_medications = None
    project_patient.patient._state.fields_cache["baseline"] = baseline

    preview = build_crf_preview(project_patient)

    assert preview["patient_baseline"] == {
        "subject_id": "",
        "name_initials": "",
        "demographics": {},
        "surgery_allergy": {},
        "comorbidities": {},
        "lifestyle": {},
        "baseline_medications": {},
    }

    assert "patient_baseline.受试者编号" in preview["missing_fields"]
    assert "patient_baseline.教育年限（年）" in preview["missing_fields"]


@pytest.mark.django_db
def test_crf_preview_preserves_non_dict_demographics_for_review(project_patient, doctor):
    baseline, _ = PatientBaseline.objects.get_or_create(
        patient=project_patient.patient,
        defaults={"created_by": doctor, "updated_by": doctor},
    )

    baseline.subject_id = "S-0001"
    baseline.demographics = "bad-demographics"
    project_patient.patient._state.fields_cache["baseline"] = baseline

    preview = build_crf_preview(project_patient)

    assert preview["patient_baseline"]["demographics"] == "bad-demographics"
    assert "patient_baseline.教育年限（年）" in preview["missing_fields"]


def evaluate(baseline=None, visits=None, registry=None):
    from apps.crf.completeness import evaluate_completeness
    from apps.crf.registry_loader import load_crf_registry

    return evaluate_completeness(registry or load_crf_registry(), baseline or {}, visits or {})


def projected(result, storage, scope="T0"):
    return next(
        f
        for s in result["sections"]
        if s["scope"] == scope
        for f in s["fields"]
        if f["storage"] == storage
    )


from tests.fixtures.crf_sentinel_cases import complete_known_case  # noqa: E402


@pytest.mark.django_db
def test_old_eleven_summary_and_completed_visits_are_not_complete(project_patient, doctor):
    PatientBaseline.objects.create(
        patient=project_patient.patient,
        subject_id="S1",
        demographics={"education_years": 9},
        created_by=doctor,
        updated_by=doctor,
    )
    for v in VisitRecord.objects.filter(project_patient=project_patient):
        v.visit_date = "2026-10-08"
        v.status = "completed"
        v.form_data = {"assessments": {"sppb": {"total": 9}, "moca": {"total": 22}}}
        v.save()
    result = build_crf_preview(project_patient)
    assert result["missing_fields"]
    assert result["sections"]
    assert result["counts"]["unresolved"] > 11


def test_all_applicable_known_values_are_complete_without_fabricated_remarks():
    baseline, visits = complete_known_case()
    result = evaluate(baseline, visits)
    assert result["missing_fields"] == []
    assert result["issues"] == []
    assert result["counts"]["normal"] > 200
    assert result["counts"]["unknown"] == 0
    assert result["counts"]["not_applicable"] == 0
    assert result["counts"]["unresolved"] == 0


@pytest.mark.parametrize("state", ["unknown", "not_applicable"])
def test_handled_state_allows_optional_reason_and_does_not_use_residual_value(state):
    baseline, visits = complete_known_case()
    key = "patient_baseline.demographics.age_years"
    baseline["demographics"]["_crf_field_meta"] = {key: {"state": state}}
    result = evaluate(baseline, visits)
    assert not any("年龄" in item for item in result["missing_fields"])
    baseline["demographics"]["_crf_field_meta"][key]["reason"] = "受试者无法提供"
    result = evaluate(baseline, visits)
    assert result["missing_fields"] == []
    f = projected(result, key, "patient_baseline")
    assert f["state"] == state and f["value"] is None and f["raw_value"] == 0
    assert result["counts"][state] == 1


def test_partial_date_unprocessed_unknown_and_conflict_preserve_both_raw_values():
    baseline, visits = complete_known_case()
    visit = visits["T0"]
    visit["visit_date"] = None
    parts = {
        "year": {"state": "value", "value": 2026},
        "month": {"state": "unknown", "value": None},
    }
    visit["form_data"]["crf"]["meta"]["visit_date_parts"] = parts
    visit["form_data"]["crf"].setdefault("_field_meta", {})["visit.visit_date"] = {
        "state": "value",
        "reason": "无法回忆月份",
    }
    result = evaluate(baseline, visits)
    assert "T0.访视日期" in result["missing_fields"]
    parts["day"] = {"state": "value", "value": 8}
    result = evaluate(baseline, visits)
    f = projected(result, "visit.visit_date")
    assert f["state"] == "unknown" and f["value"]["year"]["value"] == 2026
    assert "T0.访视日期" not in result["missing_fields"]
    visit["visit_date"] = "2026-10-08"
    result = evaluate(baseline, visits)
    f = projected(result, "visit.visit_date")
    assert f["raw_value"] == "2026-10-08" and f["partial_date_value"] == parts
    assert f["state"] == "unresolved"
    assert any("同时" in i["message"] for i in result["issues"])


def test_parent_no_keeps_inactive_rows_without_counting_them_and_yes_checks_rows():
    baseline, visits = complete_known_case()
    baseline["surgery_allergy"]["has_surgery"] = "否"
    baseline["surgery_allergy"]["surgery_records"] = [{"name": "原手术", "date": None}]
    result = evaluate(baseline, visits)
    field = projected(
        result, "patient_baseline.surgery_allergy.surgery_records", "patient_baseline"
    )
    assert not field["applicable"] and field["raw_value"][0]["name"] == "原手术"
    assert result["missing_fields"] == []
    baseline["surgery_allergy"]["has_surgery"] = "是"
    result = evaluate(baseline, visits)
    assert any("手术日期" in m for m in result["missing_fields"])
    assert (
        projected(result, field["storage"], "patient_baseline")["rows"][0]["fields"][0]["value"]
        == "原手术"
    )


def test_legacy_text_does_not_fill_details_and_unknown_json_is_preserved():
    baseline, visits = complete_known_case()
    visits["T1"]["form_data"]["crf"]["medication_change"] = "旧总文本"
    baseline["demographics"]["custom_old"] = {"note": "未识别旧记录"}
    result = evaluate(baseline, visits)
    assert any("T1." in m for m in result["missing_fields"])
    supplements = [f for s in result["sections"] if s["supplemental"] for f in s["fields"]]
    assert any(f["raw_value"] == "旧总文本" for f in supplements)
    assert any(f["raw_value"] == "未识别旧记录" for f in supplements)


def test_terminal_record_selects_whole_latest_group_and_reports_conflicts():
    baseline, visits = complete_known_case()
    visits["T0"]["form_data"]["crf"]["completion_record"] = {
        "status": "提前退出",
        "date": "2026-01-01",
    }
    visits["T2"]["form_data"]["crf"]["completion_record"] = {"status": "完成全部随访"}
    result = evaluate(baseline, visits)
    sections = [s for s in result["sections"] if s["table_ref"] == "#T29" and not s["supplemental"]]
    assert len(sections) == 1 and sections[0]["scope"] == "T2"
    assert next(f for f in sections[0]["fields"] if f["storage"].endswith(".date"))["value"] is None
    assert any(i["code"] == "conflicting_records" for i in result["issues"])
    assert any("完成/退出日期" in m for m in result["missing_fields"])


def test_same_values_keep_visit_source_and_missing_record_never_uses_computed():
    baseline, visits = complete_known_case()
    visits["T0"]["form_data"]["assessments"]["sppb"]["total"] = None
    visits["T0"]["form_data"]["computed_assessments"] = {"sppb": {"total": 12}}
    del visits["T2"]
    result = evaluate(baseline, visits)
    assert "T2.访视记录" in result["missing_fields"] and "T0.SPPB总分" in result["missing_fields"]
    for vt in ("T0", "T1"):
        field = projected(result, "visit.form_data.assessments.moca.total", vt)
        assert field["value"] == 1
        assert field["source"]["visit_type"] == vt and field["source"]["kind"] == "manual"


def test_invalid_old_numeric_and_legacy_option_are_review_issues_not_normal_answers():
    baseline, visits = complete_known_case()
    baseline["demographics"]["age_years"] = "旧的70岁"
    baseline["surgery_allergy"]["has_surgery"] = "不详"
    result = evaluate(baseline, visits)
    assert any("年龄" in m for m in result["missing_fields"])
    assert any("手术" in m for m in result["missing_fields"])
    assert (
        projected(result, "patient_baseline.demographics.age_years", "patient_baseline")[
            "raw_value"
        ]
        == "旧的70岁"
    )
    assert result["issues"]


def test_unknown_parent_residual_raw_does_not_activate_children_or_require_visible_optional_reason():
    baseline, visits = complete_known_case()
    baseline["surgery_allergy"]["has_surgery"] = "是"
    baseline["surgery_allergy"]["surgery_records"] = []
    baseline["demographics"]["_crf_field_meta"] = {
        "patient_baseline.surgery_allergy.has_surgery": {"state": "unknown", "reason": "无法确认"}
    }
    visits["T1"]["form_data"]["crf"]["adherence"]["completion_rate"] = 100
    visits["T1"]["form_data"]["crf"]["adherence"].pop("incomplete_reason", None)
    visits["T1"]["form_data"]["crf"]["_field_meta"] = {
        "visit.form_data.crf.adherence.completion_rate": {"state": "unknown", "reason": "记录缺失"}
    }
    result = evaluate(baseline, visits)
    assert result["missing_fields"] == []
    reason = projected(result, "visit.form_data.crf.adherence.incomplete_reason", "T1")
    assert reason["applicable"] and not reason["required"]
    assert not projected(
        result, "patient_baseline.surgery_allergy.surgery_records", "patient_baseline"
    )["applicable"]


def test_alternative_alcohol_units_only_require_one_handled_answer():
    baseline, visits = complete_known_case()
    baseline["lifestyle"]["drinking_status"] = "目前饮酒"
    baseline["lifestyle"]["drinking_record"]["weekly_ml"] = 0
    baseline["lifestyle"]["drinking_record"].pop("weekly_units")
    baseline["demographics"]["_crf_field_meta"] = {
        "patient_baseline.lifestyle.drinking_record.weekly_units": {"state": "value"}
    }
    result = evaluate(baseline, visits)
    assert result["missing_fields"] == []
    assert not projected(
        result, "patient_baseline.lifestyle.drinking_record.weekly_units", "patient_baseline"
    )["counted"]


def test_repeated_row_metadata_and_duplicate_ids_are_not_silently_discarded():
    baseline, visits = complete_known_case()
    baseline["surgery_allergy"]["has_surgery"] = "是"
    rows = baseline["surgery_allergy"]["surgery_records"]
    rows[0]["_field_meta"] = {"name": {"state": "unknown", "reason": "遗失"}}
    rows.append({"_row_id": "row-1", "name": "第二条", "date": "2026-01-01"})
    result = evaluate(baseline, visits)
    field = projected(
        result, "patient_baseline.surgery_allergy.surgery_records", "patient_baseline"
    )
    assert field["rows"][0]["fields"][0]["state"] == "unknown"
    assert field["rows"][1]["fields"][0]["value"] == "第二条"
    assert any(i["storage"].endswith("._row_id") for i in result["issues"])


def test_terminal_metadata_counts_as_recorded_but_empty_object_does_not():
    baseline, visits = complete_known_case()
    visits["T1"]["form_data"]["crf"]["completion_record"] = {"status": "提前退出"}
    visits["T2"]["form_data"]["crf"]["completion_record"] = {}
    result = evaluate(baseline, visits)
    assert any(
        s["scope"] == "T1" and s["table_ref"] == "#T29" and not s["supplemental"]
        for s in result["sections"]
    )
    visits["T2"]["form_data"]["crf"]["_field_meta"] = {
        "visit.form_data.crf.completion_record.status": {"state": "unknown", "reason": "待确认"}
    }
    result = evaluate(baseline, visits)
    assert any(
        s["scope"] == "T2" and s["table_ref"] == "#T29" and not s["supplemental"]
        for s in result["sections"]
    )
    assert any(
        i["code"] == "conflicting_records"
        and [v["visit_type"] for v in i["sources"]] == ["T1", "T2"]
        for i in result["issues"]
    )


def test_complete_projection_covers_all_definitions_and_all_rows_without_writes():
    from copy import deepcopy
    from apps.crf.registry_loader import load_crf_registry

    baseline, visits = complete_known_case()
    before = deepcopy((baseline, visits))
    result = evaluate(baseline, visits)
    projected_ids = {
        f["field_id"] for s in result["sections"] if not s["supplemental"] for f in s["fields"]
    }
    expected = {f["field_id"] for f in load_crf_registry()["fields"] if not f.get("legacy")}
    assert projected_ids == expected
    assert (baseline, visits) == before


@pytest.mark.django_db
def test_preview_projects_static_flow_from_controlled_original_without_answers(project_patient):
    from django.conf import settings
    from docx import Document

    preview = build_crf_preview(project_patient)
    flow = next(s for s in preview["sections"] if s["table_ref"] == "#T2")
    source = [
        [cell.text for cell in row.cells]
        for row in Document(settings.CRF_TEMPLATE_PATH).tables[2].rows
    ]
    assert len(source) == 14 and all(len(row) == 5 for row in source)
    assert flow["static_table"] == {"headers": source[0], "rows": source[1:]}
    assert flow["fields"] == [] and sum(flow["counts"].values()) == 0
    assert not any("研究流程" in label for label in preview["missing_fields"])


def test_cleared_optional_remark_with_value_metadata_stays_optional():
    baseline, visits = complete_known_case()
    baseline["comorbidities"]["other_history"] = None
    baseline["demographics"]["_crf_field_meta"] = {
        "patient_baseline.comorbidities.other_history": {"state": "value"}
    }
    result = evaluate(baseline, visits)
    assert result["missing_fields"] == []
    assert not projected(
        result, "patient_baseline.comorbidities.other_history", "patient_baseline"
    )["counted"]


def test_repeated_row_unrecognized_metadata_leaves_remain_visible_without_clinical_counts():
    from copy import deepcopy

    baseline, visits = complete_known_case()
    baseline["surgery_allergy"]["has_surgery"] = "是"
    before = evaluate(baseline, visits)
    row = baseline["surgery_allergy"]["surgery_records"][0]
    row["_field_meta"] = {
        "name": {
            "state": "value",
            "reason": "已核实",
            "old_feedback": {"note": "已知字段的旧反馈"},
        },
        "old_reason": {"state": "unknown", "reason": "旧反馈须保留"},
        "empty_old": {"note": " "},
    }
    original = deepcopy(row)
    result = evaluate(baseline, visits)
    supplemental = [f for s in result["sections"] if s["supplemental"] for f in s["fields"]]
    by_path = {f["storage"]: f for f in supplemental}
    prefix = "patient_baseline.surgery_allergy.surgery_records.0._field_meta."
    for path, value in {
        "name.old_feedback.note": "已知字段的旧反馈",
        "old_reason.state": "unknown",
        "old_reason.reason": "旧反馈须保留",
    }.items():
        item = by_path[prefix + path]
        assert item["value"] == value and item["raw_value"] == value
        assert not item["counted"] and not item["required"]
        assert "第1行" in item["label"] and "_field_meta" not in item["label"]
        assert item["source"]["kind"] == "legacy"
    assert prefix + "name.state" not in by_path
    assert prefix + "name.reason" not in by_path
    assert prefix + "empty_old.note" not in by_path
    assert (
        result["counts"] == before["counts"]
        and result["missing_fields"] == before["missing_fields"]
    )
    assert row == original


def test_repeated_row_malformed_known_metadata_leaf_is_preserved_for_review():
    baseline, visits = complete_known_case()
    baseline["surgery_allergy"]["surgery_records"][0]["_field_meta"] = {
        "name": {"state": {"old": "旧状态原文"}, "reason": ["旧原因原文"]}
    }
    result = evaluate(baseline, visits)
    supplemental = [f for s in result["sections"] if s["supplemental"] for f in s["fields"]]
    assert any(f["raw_value"] == "旧状态原文" for f in supplemental)
    assert any(f["raw_value"] == "旧原因原文" for f in supplemental)


def test_baseline_repeated_child_source_uses_full_path():
    baseline, visits = complete_known_case()
    result = evaluate(baseline, visits)
    field = projected(
        result, "patient_baseline.surgery_allergy.surgery_records", "patient_baseline"
    )
    assert field["rows"][0]["fields"][0]["source"] == {
        "kind": "patient_baseline",
        "visit_type": None,
    }
    field = projected(result, "visit.form_data.crf.medication_change.entries", "T1")
    assert field["rows"][0]["fields"][0]["source"] == {"kind": "manual", "visit_type": "T1"}

@pytest.mark.django_db
@pytest.mark.parametrize(
    ("parts", "metadata", "state", "paper_date"),
    [
        ({"year": {"state": "value", "value": 2026}, "month": {"state": "value", "value": 2}, "day": {"state": "value", "value": 28}}, None, "normal", "2026年02月28日"),
        ({"year": {"state": "value", "value": 2026}, "month": {"state": "unknown"}, "day": {"state": "unknown"}}, {"state": "value", "reason": "月份和日期记不清"}, "unknown", "2026年99月99日"),
        (None, None, "unresolved", None),
    ],
)
def test_orm_missing_native_date_preserves_parts_in_preview_and_original_docx(
    project_patient, tmp_path, parts, metadata, state, paper_date
):
    from docx import Document
    from apps.crf.services.export_docx import export_preview_to_docx

    visit = VisitRecord.objects.get(project_patient=project_patient, visit_type="T1")
    visit.visit_date = None
    visit.form_data = {"crf": {"meta": {"visit_date_parts": parts}}}
    if metadata:
        visit.form_data["crf"]["_field_meta"] = {"visit.visit_date": metadata}
    visit.save(update_fields=["visit_date", "form_data"])
    preview = build_crf_preview(project_patient)
    field = projected(preview, "visit.visit_date", "T1")
    assert field["state"] == state
    assert field["raw_value"] is None
    assert not [issue for issue in preview["issues"] if issue["scope"] == "T1" and issue["storage"] == "visit.visit_date"]
    assert ("T1.访视日期" in preview["missing_fields"]) == (state == "unresolved")
    document = Document(export_preview_to_docx(preview, tmp_path / "filled.docx"))
    date_cell = document.tables[15].rows[1].cells[1].text
    if paper_date:
        assert paper_date in date_cell
    else:
        assert "2026" not in date_cell and "99" not in date_cell


@pytest.mark.django_db
def test_project_cover_uses_current_binding_and_preserves_old_patient_document(doctor,project_patient):
    from datetime import datetime,timezone
    from apps.patients.models import PatientBaseline
    from apps.studies.models import ProjectPatient,StudyProject,StudyGroup
    from apps.crf.services.aggregate import build_crf_preview
    from apps.visits.models import VisitRecord
    baseline,_=PatientBaseline.objects.get_or_create(patient=project_patient.patient)
    old={'protocol_number':'旧方案','revision_1_0_name':'历史确认姓名'}
    baseline.demographics={'document':old};baseline.save()
    ProjectPatient.objects.filter(pk=project_patient.pk).update(enrolled_at=datetime(2026,10,9,17,30,tzinfo=timezone.utc));project_patient.refresh_from_db()
    preview=build_crf_preview(project_patient)
    projected=[f for s in preview['sections'] for f in s['fields']]
    cover=next(f for f in projected if f['field_id']=='doc_enrolled_on')
    assert cover['value']=='2026-10-10' and cover['storage']=='visit.form_data.crf.cover.enrolled_on'
    assert not any(f['field_id'].startswith('doc_revision_') for f in projected)
    assert preview['patient_baseline']['demographics']['document']==old
    other=StudyProject.objects.create(name='第二项目');group=StudyGroup.objects.create(project=other,name='第二组')
    ProjectPatient.objects.create(patient=project_patient.patient,project=other,group=group)
    preview=build_crf_preview(project_patient)
    protocol=next(f for s in preview['sections'] for f in s['fields'] if f['field_id']=='doc_protocol_number')
    assert protocol['value'] in (None,'')
    visit=VisitRecord.objects.get(project_patient=project_patient,visit_type='T0')
    visit.form_data={'crf':{'cover':{'protocol_number':'本项目','enrolled_on':None}}};visit.save()
    preview=build_crf_preview(project_patient)
    assert next(f for s in preview['sections'] for f in s['fields'] if f['field_id']=='doc_protocol_number')['value']=='本项目'
    assert next(f for s in preview['sections'] for f in s['fields'] if f['field_id']=='doc_enrolled_on')['value'] is None


@pytest.mark.django_db
@pytest.mark.parametrize('state',['unknown','not_applicable'])
def test_old_cover_state_and_reason_survive_readonly_projection(project_patient,state):
    from apps.patients.models import PatientBaseline
    baseline,_=PatientBaseline.objects.get_or_create(patient=project_patient.patient)
    cover={'protocol_number':'旧方案','enrolled_on':None,'investigator_name':'旧医生'}
    baseline.demographics={'document':cover,'_crf_field_meta':{f'patient_baseline.demographics.document.{key}':{'state':state,'reason':'原理由','extra':'原补充'} for key in cover}}
    baseline.save()
    before=deepcopy(baseline.demographics)
    preview=build_crf_preview(project_patient)
    for key in cover:
        item=projected(preview,f'visit.form_data.crf.cover.{key}','T0')
        assert item['state']==state
        assert item['raw_value']=={'__crf_answer__':1,'state':state,'value':cover[key],'reason':'原理由'}
    baseline.refresh_from_db();assert baseline.demographics==before

@pytest.mark.django_db
def test_old_client_complete_parts_roundtrip_preview_and_docx(doctor,project_patient,tmp_path):
    from rest_framework.test import APIClient
    from docx import Document
    from apps.crf.services.export_docx import export_preview_to_docx
    parts={'year':{'state':'value','value':2026},'month':{'state':'value','value':2},'day':{'state':'value','value':28}}
    visit=VisitRecord.objects.get(project_patient=project_patient,visit_type='T1')
    c=APIClient();c.force_authenticate(doctor)
    url=f'/api/visits/{visit.pk}/'
    response=c.patch(url,{'visit_date':None,'form_data':{'crf':{'meta':{'visit_date_parts':parts}}}},format='json')
    assert response.status_code==200,response.data
    assert c.get(url).data['visit_date']==parts
    preview=build_crf_preview(project_patient)
    item=projected(preview,'visit.visit_date','T1')
    assert item['state']=='normal',item
    assert not item['issues']
    doc=Document(export_preview_to_docx(preview,tmp_path/'parts.docx'))
    assert '2026年02月28日' in doc.tables[15].rows[1].cells[1].text
