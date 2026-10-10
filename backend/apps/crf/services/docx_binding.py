"""Fill the controlled paper form from the shared read-only clinical projection.

Coordinates are physical w:tc indices. Word bookmarks preserve original row
identity across repeat-row expansion and a save/reopen cycle.
"""

from collections import defaultdict
from copy import deepcopy
import json
import re

from docx.document import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell

from apps.crf.registry_loader import CrfAssetError, load_crf_source_slots


_DATE_BLANK = r"(?:\|__){4}\|年(?:\|__){2}\|月(?:\|__){2}\|日|_{4}年_{2,4}月(?:_{2}日)?"
_BLANK = re.compile(_DATE_BLANK + r"|_+|□{4}")
# Existing remark columns in the hash-pinned paper form: table, first/last
# original row, physical remark cell. Satisfaction row 6 is a total, not a note.
_REMARK_REGIONS = (
    (5, 1, 4, 4),
    (6, 1, 6, 4),
    (9, 2, 4, 2),
    (9, 7, 8, 3),
    (10, 1, 14, 4),
    (19, 1, 5, 6),
    (20, 2, 3, 5),
    (26, 1, 5, 6),
    (27, 2, 3, 5),
    (30, 1, 6, 4),
)


def _remark_position(position):
    for table, first, last, column in _REMARK_REGIONS:
        if position["table_index"] == table and first <= position["row_index"] <= last:
            return {**position, "cell_index": column, "cell_indices": [column]}
    return None


def _row_name(table_index, row_index):
    return f"CRF_T{table_index:02d}_R{row_index:02d}"


def _source_row(document, position):
    if position.get("index_space") != "xml_tc":
        raise ValueError("CRF 位置必须声明 index_space=xml_tc")
    table = document.tables[position["table_index"]]
    name = position.get("row_anchor") or _row_name(position["table_index"], position["row_index"])
    for row in table._tbl.tr_lst:
        if any(b.get(qn("w:name")) == name for b in row.iter(qn("w:bookmarkStart"))):
            return table, row
    if position.get("row_anchor") or any(
        b.get(qn("w:name"), "").startswith("CRF_T") for b in table._tbl.iter(qn("w:bookmarkStart"))
    ):
        raise ValueError(f"CRF 原行锚点缺失：{name}")
    return table, table._tbl.tr_lst[position["row_index"]]


def resolve_source_cells(document: Document, position: dict) -> list[_Cell]:
    """Resolve original xml_tc coordinates, including explicit multi-cell choices."""
    table, row = _source_row(document, position)
    return [
        _Cell(row.tc_lst[i], table) for i in position.get("cell_indices", [position["cell_index"]])
    ]


def _bookmark(row, name, identifier):
    paragraph = row.tc_lst[0].p_lst[0]
    start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name)
    end.set(qn("w:id"), str(identifier))
    insertion = 1 if paragraph.pPr is not None else 0
    paragraph.insert(insertion, start)
    paragraph.insert(insertion + 1, end)


def _recipes(manifest):
    """Pair ordered semantic slots with their reviewed original cell contents.

    Each cursor is confined to one physical cell. Options come from the reviewed
    source manifest, including print-only choices; no field-label search occurs.
    """
    cursors = defaultdict(int)
    result = {}
    for slot in manifest["slots"]:
        key = (slot["table_index"], slot["row_index"], slot["cell_index"], bool(slot["options"]))
        text = slot["source_cells"][0]
        recipe = {"slot": slot, "choices": []}
        if len(slot["cell_indices"]) > 1:
            recipe["choices"] = [
                (str(option), index, 0, 1) for index, option in enumerate(slot["options"])
            ]
        elif slot["options"]:
            cursor = cursors[key]
            for option in slot["options"]:
                index = text.find("□" + option, cursor)
                if index < 0:
                    raise CrfAssetError(f"CRF 选项位置不匹配：{slot['slot_id']} {option}")
                recipe["choices"].append((option, 0, index, index + 1))
                cursor = index + 1 + len(option)
            cursors[key] = cursor
            extra = _BLANK.search(text, cursor)
            if extra:
                recipe["other_span"] = extra.span()
        else:
            match = _BLANK.search(text, cursors[key])
            if match:
                recipe["span"] = match.span()
                cursors[key] = match.end()
            elif not text:
                recipe["span"] = (0, 0)
            else:
                raise CrfAssetError(f"CRF 文字填写位不匹配：{slot['slot_id']}")
        result[slot["slot_id"]] = recipe
    return result


def _date_text(value, definition, state):
    components = (
        ("year", "month")
        if definition.get("date_precision") == "month"
        else ("year", "month", "day")
    )
    units = {"year": "年", "month": "月", "day": "日"}
    if isinstance(value, dict):
        parts = []
        for key in components:
            component = value.get(key)
            component = component if isinstance(component, dict) else {}
            if component.get("state") == "unknown":
                item = "9999" if key == "year" else "99"
            elif component.get("state") == "value" and component.get("value") is not None:
                item = str(component["value"]).zfill(4 if key == "year" else 2)
            else:
                item = "____" if key == "year" else "__"
            parts.append(item + units[key])
        return "".join(parts)
    if state == "unknown":
        return "".join(("9999" if key == "year" else "99") + units[key] for key in components)
    if state == "not_applicable":
        return "88"
    return "".join(part + units[key] for key, part in zip(components, str(value).split("-")))


def _effective_state(field):
    metadata = field.get("metadata")
    if (
        field["state"] == "unresolved"
        and not field.get("issues")
        and isinstance(metadata, dict)
    ):
        if metadata.get("state") in ("unknown", "not_applicable"):
            return metadata["state"]
    return field["state"]


def _display(field, definition):
    state = _effective_state(field)
    if not field.get("applicable") or state == "unresolved":
        return None
    value = field.get("value")
    if definition["widget"] == "date":
        return _date_text(value, definition, state)
    if state in ("unknown", "not_applicable"):
        return "99" if state == "unknown" else "88"
    return value


def _replace_ranges(cell, replacements):
    """Replace only text ranges, retaining paragraph/run and cell formatting."""
    if not replacements:
        return
    # Controlled source cells have a single paragraph; run boundaries can split
    # underscores and option text, so edit the underlying text nodes by offsets.
    nodes = list(cell._tc.iter(qn("w:t")))
    for start, end, value in sorted(
        replacements, key=lambda item: (item[0], item[1]), reverse=True
    ):
        if not nodes:
            cell.paragraphs[0].add_run(value)
            nodes = list(cell._tc.iter(qn("w:t")))
            continue
        offset = 0
        overlaps = []
        for node in nodes:
            length = len(node.text or "")
            if offset <= start <= offset + length and not overlaps:
                overlaps.append((node, start - offset, min(end - offset, length)))
            elif offset < end and offset + length > start:
                overlaps.append((node, 0, min(end - offset, length)))
            offset += length
        if not overlaps:
            raise CrfAssetError("CRF 文字区间超出原格")
        for index, (node, left, right) in enumerate(overlaps):
            node.text = (
                (node.text or "")[:left] + (value if index == 0 else "") + (node.text or "")[right:]
            )
            node.set(qn("xml:space"), "preserve")
    # Word uses explicit line-break elements; literal newlines in w:t are not
    # enough for interoperable PDF rendering of multiline feedback.
    for node in list(cell._tc.iter(qn("w:t"))):
        parts = (node.text or "").split("\n")
        if len(parts) < 2:
            continue
        node.text = parts[0]
        cursor = node
        for part in parts[1:]:
            line_break, text_node = OxmlElement("w:br"), OxmlElement("w:t")
            text_node.text = part
            text_node.set(qn("xml:space"), "preserve")
            cursor.addnext(line_break)
            line_break.addnext(text_node)
            cursor = text_node


def _complete_repeat_header_prefix(document: Document, registry: dict) -> None:
    """Keep the source's introductory row with its repeated table header.

    Linux LibreOffice can lose the remaining rows when a table's first row is
    ordinary and the following row is a repeat header. Only fill missing header
    flags before the first existing header and before any bound data row. Tables
    with multiple clinical repeat sections retain their original header flags;
    otherwise a surgery header would incorrectly repeat above allergy records.
    """
    repeat_regions = defaultdict(set)
    for field in registry["fields"]:
        if field.get("widget") != "table_repeat":
            continue
        first_data_rows = {}
        for child in field["children"]:
            for position in child["source_positions"]:
                index, row = position["table_index"], position["row_index"]
                first_data_rows[index] = min(first_data_rows.get(index, row), row)
        for index, first_data in first_data_rows.items():
            repeat_regions[index].add(first_data)
    for index, starts in repeat_regions.items():
        if len(starts) != 1:
            continue
        first_data = next(iter(starts))
        rows = document.tables[index]._tbl.tr_lst
        first_header = next(
            (i for i, row in enumerate(rows[:first_data])
             if row.trPr is not None and row.trPr.find(qn("w:tblHeader")) is not None),
            None,
        )
        if first_header is not None:
            for row in rows[:first_header]:
                row.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def fill_crf_document(document: Document, preview: dict, registry: dict) -> Document:
    _complete_repeat_header_prefix(document, registry)
    manifest = load_crf_source_slots()
    recipes = _recipes(manifest)
    definitions = {}
    for definition in registry["fields"]:
        definitions[definition["field_id"]] = definition
        definitions.update({child["field_id"]: child for child in definition.get("children", [])})
    existing_ids = [
        int(b.get(qn("w:id")))
        for b in document.element.iter(qn("w:bookmarkStart"))
        if b.get(qn("w:id"), "").isdigit()
    ]
    bookmark_id = max(existing_ids, default=0) + 1
    for ti, table in enumerate(document.tables):
        for ri, row in enumerate(table._tbl.tr_lst):
            _bookmark(row, _row_name(ti, ri), bookmark_id)
            bookmark_id += 1

    replacements = defaultdict(list)
    cells = {}
    supplements = []
    row_notes = defaultdict(list)

    def note(section, field, text):
        if text:
            label = field["label"]
            row = re.search(r"\.(\d+)(?:\.|$)", field.get("storage", ""))
            if row and not re.search(r"第\d+行", label):
                label = f"第{int(row[1]) + 1}行 · {label}"
            title = registry.get("table_titles", {}).get(section.get("table_ref"), section["title"])
            supplements.append((section["scope"], title, label, text))

    def preserve_inactive(section, field):
        parts = []
        state = _effective_state(field)
        if state in ("unknown", "not_applicable"):
            parts.append("原状态：" + ("不详" if state == "unknown" else "不适用"))
        if field.get("reason"):
            parts.append("原说明：" + field["reason"])
        for key, label in (("raw_value", "原值"), ("partial_date_value", "原日期分量")):
            if field.get(key) not in (None, "", [], {}):
                parts.append(label + "：" + json.dumps(field[key], ensure_ascii=False))
        if parts:
            note(section, field, "历史非活动记录（不作为当前答案）；" + "；".join(parts))

    def fill(section, field, position, definition):
        if not field.get("applicable"):
            return
        recipe = recipes[position["slot_id"]]
        targets = resolve_source_cells(document, position)
        for target in targets:
            cells[target._tc] = target
        state = _effective_state(field)
        reason = field.get("reason")
        if reason or state in ("unknown", "not_applicable"):
            status = {"unknown": "99 不详", "not_applicable": "88 不适用"}.get(state, "")
            detail = status + ("；" if status and reason else "") + (reason or "")
            text = f"{field['label']}：{detail}"
            remark_position = _remark_position(position)
            if remark_position:
                remark = resolve_source_cells(document, remark_position)[0]
                cells[remark._tc] = remark
                row_notes[remark._tc].append(text)
            else:
                note(section, field, text)
        value = _display(field, definition)
        if value is None or value == "":
            return
        if recipe["choices"]:
            mapped = definition.get("export_choice_mapping", {}).get(str(value), {})
            selected = mapped.get("value", value)
            if state in ("unknown", "not_applicable"):
                selected = definition.get("metadata_choice_mapping", {}).get(state)
                if selected is None:
                    # Status is not a clinical choice. Leave all printed boxes empty.
                    targets[0].paragraphs[0].add_run(f" [{field['label']}：{value}]")
            selected = selected if isinstance(selected, list) else [selected]
            for option, ci, start, end in recipe["choices"]:
                if option in [str(item) for item in selected]:
                    replacements[targets[ci]._tc].append((start, end, "☑"))
            if mapped.get("other_text") and recipe.get("other_span"):
                replacements[targets[0]._tc].append((*recipe["other_span"], mapped["other_text"]))
        else:
            combination = definition.get("source_combination")
            rendered = str(value)
            if combination:
                rendered = combination["label"] + "：" + rendered
            replacements[targets[0]._tc].append(
                (*recipe["span"], rendered, combination.get("order", 0) if combination else 0)
            )

    for section in preview.get("sections", []):
        for field in section.get("fields", []):
            definition = definitions.get(field["field_id"])
            if (
                section.get("supplemental")
                or not definition
                or field.get("source", {}).get("kind") == "legacy"
            ):
                raw = field.get("raw_value")
                if raw is not None:
                    note(
                        section, field, "历史补充（需核对）：" + json.dumps(raw, ensure_ascii=False)
                    )
                continue
            if not field.get("applicable"):
                preserve_inactive(section, field)
                continue
            if not field.get("source_positions"):
                state = _effective_state(field)
                if state in ("unknown", "not_applicable"):
                    code = "99 不详" if state == "unknown" else "88 不适用"
                    note(section, field, code + ("；" + field["reason"] if field.get("reason") else ""))
                elif field.get("reason"):
                    note(section, field, "说明：" + field["reason"])
            if field.get("state") == "unresolved" and field.get("raw_value") not in (
                None,
                "",
                [],
                {},
            ):
                note(
                    section,
                    field,
                    "待核对原值：" + json.dumps(field["raw_value"], ensure_ascii=False),
                )
            if field["state"] == "unresolved" and field.get("partial_date_value") is not None:
                note(
                    section,
                    field,
                    "待核对日期分量原值："
                    + json.dumps(field["partial_date_value"], ensure_ascii=False),
                )
            if definition["widget"] != "table_repeat":
                for position in field.get("source_positions", []):
                    fill(section, field, position, definition)
                continue
            if not field.get("applicable") or _effective_state(field) in (
                "unknown",
                "not_applicable",
            ):
                continue
            child_definitions = definition["children"]
            positions = child_definitions[0]["source_positions"]
            capacity = len(positions)
            table, final_row = _source_row(document, positions[-1])
            pristine_row = deepcopy(final_row)
            for row in field["rows"]:
                row_index = row["index"]
                anchor = None
                if row_index >= capacity:
                    new_row = deepcopy(pristine_row)
                    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
                        for bookmark in list(new_row.iter(qn(tag))):
                            bookmark.getparent().remove(bookmark)
                    anchor = f"CRF_X{bookmark_id}"
                    _bookmark(new_row, anchor, bookmark_id)
                    bookmark_id += 1
                    final_row.addnext(new_row)
                    final_row = new_row
                    if positions[0]["table_index"] in (21, 28):
                        _replace_ranges(
                            _Cell(new_row.tc_lst[0], table),
                            [(0, len(_Cell(new_row.tc_lst[0], table).text), str(row_index + 1))],
                        )
                for child in row["fields"]:
                    if not child.get("applicable"):
                        preserve_inactive(section, child)
                        continue
                    child_definition = definitions[child["field_id"]]
                    source = child_definition["source_positions"][min(row_index, capacity - 1)]
                    position = dict(source)
                    if anchor:
                        position["row_anchor"] = anchor
                    if child["state"] == "unresolved" and child.get("raw_value") not in (
                        None,
                        "",
                        [],
                        {},
                    ):
                        note(
                            section,
                            child,
                            "待核对原值：" + json.dumps(child["raw_value"], ensure_ascii=False),
                        )
                    fill(section, child, position, child_definition)

    for tc, edits in replacements.items():
        # Explicit registry join_labeled members share one source blank.
        combined = defaultdict(list)
        for edit in edits:
            start, end, value = edit[:3]
            combined[start, end].append((edit[3] if len(edit) > 3 else 0, value))
        _replace_ranges(
            cells[tc],
            [
                (start, end, "\n".join(value for _, value in sorted(values)))
                for (start, end), values in combined.items()
            ],
        )
    for tc, notes in row_notes.items():
        cell = cells[tc]
        cell.paragraphs[0].add_run(("\n" if cell.text else "") + "\n".join(notes))
    for issue in preview.get("issues", []):
        if issue["code"] == "conflicting_records":
            supplements.append(
                (
                    issue["scope"],
                    "记录冲突需核对",
                    "终末记录冲突",
                    issue["message"]
                    + "\n各访视原记录："
                    + json.dumps(issue.get("sources", []), ensure_ascii=False),
                )
            )
    if supplements:
        document.add_page_break()
        document.add_heading("填报说明与历史补充", level=1)
        current = None
        scope_titles = {
            "patient_baseline": "患者基线",
            "T0": "T0 基线访视",
            "T1": "T1 随访",
            "T2": "T2 随访",
            "study": "研究记录",
        }
        for scope, title, label, text in supplements:
            if (scope, title) != current:
                document.add_heading(f"{scope_titles.get(scope, '研究记录')} {title}", level=2)
                current = (scope, title)
            document.add_paragraph(f"{label}：{text}")
    return document
